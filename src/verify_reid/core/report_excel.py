"""Excel レポート生成（openpyxl）。シート: 条件 / 外見比較 / 外見比較(追跡ID集約) / ログ。"""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import cv2
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .gallery import GalleryEntry
from .phase_log import PhaseLogger
from .similarity import Pair, competition_ranks

GRID_COLS = 10
COL_PX = 165
INFO_ROW_PX = 120
IMG_ROW_PX = 200
IMG_PAD = 4
QUERY_FILL = PatternFill("solid", fgColor="FFD9EAD3")

MODE_IMAGE = "外見画像指定"
MODE_TRACK = "追跡ID指定"

_SKELETON = [(0, 1), (0, 2), (1, 3), (2, 4), (5, 6), (5, 7), (7, 9), (6, 8), (8, 10),
             (5, 11), (6, 12), (11, 12), (11, 13), (13, 15), (12, 14), (14, 16)]


@dataclass
class ReportConditions:
    model_name: str
    mode: str
    query_path: str
    gallery_dir: str
    output_dir: str
    skeleton: bool
    generated_at: datetime


def _px_to_width(px: int) -> float:
    return max(1.0, (px - 5) / 7)


def _px_to_pt(px: int) -> float:
    return px * 0.75


def _fmt_conf(c) -> str:
    return f"{c:.3f}" if isinstance(c, (int, float)) else "-"


def _info_text(e: GalleryEntry, head: str, sim: float | None = None, extra: str = "") -> str:
    lines = [head + (f" sim={sim:.4f}" if sim is not None else "")]
    if extra:
        lines.append(extra)
    lines += [
        f"Loc:{e.location} / Cam:{e.camera_id}",
        f"Track:{e.track_id} / 信頼度:{_fmt_conf(e.person_conf)}",
        f"映像:{e.video_file}",
        e.captured_text,
    ]
    return "\n".join(lines)


class _BytesImage(XLImage):
    """バイト列を保持する画像。保存で元ストリームが閉じられても再保存できる。"""

    def __init__(self, raw: bytes):
        self._raw = raw
        super().__init__(io.BytesIO(raw))

    def _data(self):
        return self._raw


def _skeleton_png(e: GalleryEntry) -> bytes | None:
    """json の骨格(フレーム座標)を _body.jpg 上に描画して PNG にする。"""
    if not e.keypoints or not e.bbox_xyxy:
        return None
    img = cv2.imread(str(e.body_path))
    if img is None:
        return None
    x0, y0 = e.bbox_xyxy[0], e.bbox_xyxy[1]
    pts = [(k["x"] - x0, k["y"] - y0, k["conf"]) for k in e.keypoints]
    for a, b in _SKELETON:
        if pts[a][2] >= 0.3 and pts[b][2] >= 0.3:
            cv2.line(img, (int(pts[a][0]), int(pts[a][1])), (int(pts[b][0]), int(pts[b][1])), (0, 255, 0), 2)
    for x, y, c in pts:
        if c >= 0.3:
            cv2.circle(img, (int(x), int(y)), 3, (0, 0, 255), -1)
    ok, buf = cv2.imencode(".png", img)
    return buf.tobytes() if ok else None


def _place_image(ws, cell: str, e: GalleryEntry, skeleton: bool = False) -> None:
    src = _skeleton_png(e) if skeleton else None
    try:
        img = _BytesImage(src) if src is not None else XLImage(str(e.body_path))
    except Exception:
        return
    max_w, max_h = COL_PX - IMG_PAD * 2, IMG_ROW_PX - IMG_PAD * 2
    scale = min(max_w / img.width, max_h / img.height, 1.0)  # 縮小のみ
    img.width, img.height = max(1, int(img.width * scale)), max(1, int(img.height * scale))
    ws.add_image(img, cell)


def _setup_grid(ws) -> None:
    for c in range(1, GRID_COLS + 1):
        ws.column_dimensions[get_column_letter(c)].width = _px_to_width(COL_PX)


def _write_cell(ws, start_row: int, pos: int, e: GalleryEntry, text: str,
                skeleton: bool = False, fill: PatternFill | None = None) -> None:
    """pos 番目(0始まり)のセルを 2行1組(情報行+画像行)で書き込む。"""
    row = start_row + (pos // GRID_COLS) * 2
    col = pos % GRID_COLS + 1
    info = ws.cell(row=row, column=col, value=text)
    info.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    info.font = Font(size=9)
    if fill:
        info.fill = fill
    ws.row_dimensions[row].height = _px_to_pt(INFO_ROW_PX)
    ws.row_dimensions[row + 1].height = _px_to_pt(IMG_ROW_PX)
    _place_image(ws, f"{get_column_letter(col)}{row + 1}", e, skeleton)


def _grid_rows(n: int) -> int:
    return ((n + GRID_COLS - 1) // GRID_COLS) * 2


def _build_conditions(ws, cond: ReportConditions, n_query: int, n_gallery: int,
                      n_tracks: int, n_out: int) -> None:
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 90
    query_label = "クエリ外見画像" if cond.mode == MODE_IMAGE else "クエリ追跡IDフォルダ"
    rows = [
        ("生成日時", cond.generated_at.strftime("%Y/%m/%d %H:%M:%S")),
        ("特徴量抽出モデル", cond.model_name),
        ("クエリ指定方法", cond.mode),
        (query_label, cond.query_path),
        ("比較対象フォルダ", cond.gallery_dir),
        ("レポート出力先フォルダ", cond.output_dir),
        ("骨格描画", "ON" if cond.skeleton else "OFF"),
        ("クエリ画像数", n_query),
        ("比較対象 画像数", n_gallery),
        ("比較対象 追跡ID数", n_tracks),
        ("外見比較シート 出力件数", n_out),
    ]
    for i, (k, v) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)


def _build_compare_sheet(ws, mode: str, queries: list[GalleryEntry], gallery: list[GalleryEntry],
                         pairs: list[Pair], skeleton: bool) -> None:
    _setup_grid(ws)
    for i, q in enumerate(queries):
        _write_cell(ws, 1, i, q, _info_text(q, "クエリ"), skeleton=skeleton, fill=QUERY_FILL)
    start = 1 + _grid_rows(len(queries)) + 1  # 1行空ける
    ranks = competition_ranks([p.sim for p in pairs])
    for i, (p, rank) in enumerate(zip(pairs, ranks)):
        g = gallery[p.g_idx]
        extra = f"Query:{queries[p.q_idx].captured_text}" if mode == MODE_TRACK else ""
        _write_cell(ws, start, i, g, _info_text(g, f"{rank}位", p.sim, extra))


def _build_log_sheet(ws, log: PhaseLogger, note: str = "") -> None:
    heads = ["No", "フェーズ", "開始時刻", "終了時刻", "所要時間(秒)", "備考"]
    widths = [6, 40, 24, 24, 14, 70]
    for c, (h, w) in enumerate(zip(heads, widths), start=1):
        ws.cell(row=1, column=c, value=h).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(c)].width = w

    def ms(dt: datetime) -> str:
        return dt.strftime("%Y/%m/%d %H:%M:%S.") + f"{dt.microsecond // 1000:03d}"

    row = 2
    ws.cell(row=row, column=2, value="レポート出力 全体")
    ws.cell(row=row, column=3, value=ms(log.started_at))
    if log.finished_at:
        ws.cell(row=row, column=4, value=ms(log.finished_at))
        ws.cell(row=row, column=5, value=round((log.finished_at - log.started_at).total_seconds(), 3))
    ws.cell(row=row, column=6, value=note)
    for rec in log.records:
        row += 1
        ws.cell(row=row, column=1, value=row - 2)
        ws.cell(row=row, column=2, value=rec.name)
        ws.cell(row=row, column=3, value=ms(rec.start))
        if rec.end:
            ws.cell(row=row, column=4, value=ms(rec.end))
            ws.cell(row=row, column=5, value=round(rec.elapsed_sec, 3))
        ws.cell(row=row, column=6, value=rec.note)


def build_report(path: Path, cond: ReportConditions, queries: list[GalleryEntry],
                 gallery: list[GalleryEntry], ranked: list[Pair], aggregated: list[Pair],
                 log: PhaseLogger) -> None:
    """レポートを作成して保存する。ログシートは全フェーズ終了後に追記して再保存する。"""
    wb = Workbook()
    ws_cond = wb.active
    ws_cond.title = "条件"
    n_tracks = len({g.track_key for g in gallery})
    _build_conditions(ws_cond, cond, len(queries), len(gallery), n_tracks, len(ranked))
    with log.phase("Excel作成（シート生成・画像配置）"):
        _build_compare_sheet(wb.create_sheet("外見比較"), cond.mode, queries, gallery, ranked, cond.skeleton)
        _build_compare_sheet(wb.create_sheet("外見比較(追跡ID集約)"), cond.mode, queries, gallery,
                             aggregated, cond.skeleton)
    with log.phase("Excel保存"):
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
    log.finish()
    _build_log_sheet(wb.create_sheet("ログ"), log, "ログシート追記のための再保存は記録対象外")
    wb.save(path)
