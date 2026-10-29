"""顔サイズ別精度検証の Excel 出力（dev_work\\比較結果類似度.xlsx「比較結果類似度まとめ【サンプル】」「比較結果誤識別数まとめ【サンプル】」のレイアウト）。

画像 N 枚を顔 bbox サイズ順（呼び出し側で並べ済み）に、5〜11 行目の I 列より右（横）と B〜H 列の 12 行目より下（縦）に同じ並びで置く。
類似度シート: 縦 i × 横 j のうち i > j（左下半分）だけに類似度を出し、類似度の段階（SIM_BINS）で色を付ける（右上半分は左下と同じ値なので出さない）。
誤識別数シート: 同じ枠の対角以外すべてに「i との類似度が j との類似度より高い非対象者の画像の数」を出し、件数の段階（count_bin）で色を付ける。
9 行目 / F 列の「顔の向き」は後から人手で式を入れるので値は空欄。"""
from __future__ import annotations

import math
import shutil
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import numpy as np
import cv2
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..common.imageio import imread, imwrite
from ..common.phase_log import PhaseLogger
from .excel import _build_log_sheet
from .gallery import GalleryEntry
from .query_set_excel import BORDER, CENTER, HEAD_FILL, HEAD_FONT, NA

SHEET_NAME = "比較結果類似度まとめ"
MISIDENTIFY_SHEET_NAME = "比較結果誤識別数まとめ"
INFO_SHEET_NAME = "対象画像情報"
FIRST_ROW = 12   # 縦側の最初の画像の行
FIRST_COL = 9    # 横側の最初の画像の列（I）

# 類似度の段階（下限, 色）。下限以上で最後に当てはまる段階の色を塗る。上の段階ほど似ている（赤 → 黄 → 緑）
SIM_BINS = (
    (float("-inf"), "FFF8696B"),
    (0.2, "FFFA9A72"),
    (0.3, "FFFDC17B"),
    (0.4, "FFFFEB84"),
    (0.5, "FFD1E183"),
    (0.6, "FFA3D17F"),
    (0.7, "FF63BE7B"),
)

# 誤識別の件数の段階: 0 / 1〜10 / 11〜20 / 21〜50 / 51〜100 / 以降 50 件刻み。MAX_COUNT_BIN 以上は同じ色で打ち止め
MAX_COUNT_BIN = 12   # 451 件以上
_COUNT_COLORS = ((0xFF, 0xEB, 0x84), (0xFD, 0xB0, 0x7A), (0xF8, 0x69, 0x6B), (0xC0, 0x00, 0x00))  # 黄 → 橙 → 赤 → 濃赤
_COUNT_LEGEND = (0, 1, 11, 21, 51, 101, 10_000)  # 凡例に出す段階の代表件数

# 5〜11 行目の行見出し（H 列）と、画像 1 枚の値の取り出し
_ROW_LABELS = ("カメラID", "映像ファイル名", "追跡ID", "画像ファイル名", "顔の向き", "画像サイズ(幅:px)", "画像サイズ(高さ:px)")
# 11 行目の縦側見出し（B〜H 列）と、画像 1 枚の値の取り出し
_COL_HEADS = ("カメラ\nID", "追跡\nID", "映像\nファイル名", "画像\nファイル名", "顔の向き", "画像サイズ\n(幅:px)", "画像サイズ\n(高さ:px)")


@dataclass
class VerifyConditions:
    model_name: str
    landmark_name: str
    gate_filter_text: str
    generated_at: datetime
    gallery_dir: str = ""            # 比較対象フォルダ（誤識別レポートのとき）
    n_nontarget: int | None = None   # 非対象者の画像数（誤識別レポートのとき）
    log_text_name: str = ""          # 誤認した画像の一覧（テキストファイル名。誤識別レポートのとき）


def sim_fill(sim: float) -> PatternFill:
    color = next(c for lo, c in reversed(SIM_BINS) if sim >= lo)
    return PatternFill("solid", fgColor=color)


def bin_label(i: int) -> str:
    lo = SIM_BINS[i][0]
    if i == 0:
        return f"{SIM_BINS[1][0]:.1f} 未満"
    if i == len(SIM_BINS) - 1:
        return f"{lo:.1f} 以上"
    return f"{lo:.1f} 〜 {SIM_BINS[i + 1][0]:.1f}"


def count_bin(n: int) -> int:
    """件数の段階（0: 0 件 / 1: 1〜10 / 2: 11〜20 / 3: 21〜50 / 4: 51〜100 / 5〜: 101 件以降 50 件刻み。MAX_COUNT_BIN で打ち止め）。"""
    if n <= 0:
        return 0
    for b, upper in enumerate((10, 20, 50, 100), start=1):
        if n <= upper:
            return b
    return min(MAX_COUNT_BIN, 4 + math.ceil((n - 100) / 50))


def count_bin_label(b: int) -> str:
    if b <= 4:
        return ("0 件", "1〜10 件", "11〜20 件", "21〜50 件", "51〜100 件")[b]
    if b >= MAX_COUNT_BIN:
        return f"{101 + (MAX_COUNT_BIN - 5) * 50} 件以上"
    start = 101 + (b - 5) * 50
    return f"{start}〜{start + 49} 件"


def count_fill(n: int) -> PatternFill:
    """0 件は緑。1 件以上は段階が上がるほど 黄 → 橙 → 赤 → 濃赤。"""
    b = count_bin(n)
    if b == 0:
        return PatternFill("solid", fgColor="FF63BE7B")
    t = (b - 1) / (MAX_COUNT_BIN - 1) * (len(_COUNT_COLORS) - 1)
    k = min(int(t), len(_COUNT_COLORS) - 2)
    a, c = _COUNT_COLORS[k], _COUNT_COLORS[k + 1]
    rgb = [round(x + (y - x) * (t - k)) for x, y in zip(a, c)]
    return PatternFill("solid", fgColor="FF" + "".join(f"{v:02X}" for v in rgb))


def _px(v) -> object:
    return NA if v is None else round(float(v), 1)


def _values(e: GalleryEntry) -> list:
    """横側（5〜11 行目）の画像 1 枚の値。並びは _ROW_LABELS（カメラID・映像ファイル名・追跡ID・画像ファイル名・顔の向き（空欄）・幅・高さ）。"""
    return [e.camera_id, e.video_file, e.track_id, e.face_path.name, None, _px(e.face_width), _px(e.face_height)]


def _col_values(e: GalleryEntry) -> list:
    """縦側（12 行目以降・B〜H 列）の画像 1 枚の値。並びは _COL_HEADS（カメラID・追跡ID・映像ファイル名・画像ファイル名・顔の向き（空欄）・幅・高さ）。
    横側と並びが違う（サンプルのレイアウトどおり）ので _values とは別にしている。"""
    v = _values(e)
    return [v[0], v[2], v[1], v[3], v[4], v[5], v[6]]


def _put(ws, r: int, c: int, v, fill: PatternFill | None = None, head: bool = False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.border = BORDER
    cell.alignment = CENTER
    if head:
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    elif fill is not None:
        cell.fill = fill
    return cell


def _build_matrix_sheet(ws, entries: list[GalleryEntry], title: str, note: str, legend: list[tuple[str, PatternFill]],
                        cell) -> None:
    """行列シートの共通部。cell(i, j) は縦 i × 横 j のセルに入れる (値, 塗り, 表示形式)。出さないなら None（対角は常に空欄の枠）。"""
    n = len(entries)
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    ws.cell(row=2, column=2, value=note).font = Font(size=9)
    for i, (label, fill) in enumerate(legend):
        _put(ws, 3, 2 + i, label, fill=fill).font = Font(size=9)
    # 横側（5〜11 行目）。H 列が行見出し、I 列以降が画像 j
    for k, label in enumerate(_ROW_LABELS):
        _put(ws, 5 + k, 8, label, head=True)
    for j, e in enumerate(entries):
        for k, v in enumerate(_values(e)):
            _put(ws, 5 + k, FIRST_COL + j, v)
    # 縦側（11 行目が見出し、12 行目以降が画像 i）
    for k, h in enumerate(_COL_HEADS):
        _put(ws, 11, 2 + k, h, head=True)
    for i, e in enumerate(entries):
        for k, v in enumerate(_col_values(e)):
            _put(ws, FIRST_ROW + i, 2 + k, v)
    for i in range(n):
        for j in range(n):
            if i == j:
                _put(ws, FIRST_ROW + i, FIRST_COL + j, None)  # 対角は枠だけ（自分自身）
                continue
            v = cell(i, j)
            if v is not None:
                value, fill, fmt = v
                _put(ws, FIRST_ROW + i, FIRST_COL + j, value, fill=fill).number_format = fmt
    for r in range(5, 11):
        ws.row_dimensions[r].height = 15.75
    ws.row_dimensions[8].height = 62
    ws.row_dimensions[6].height = 30
    ws.row_dimensions[11].height = 32.25
    for i in range(n):
        ws.row_dimensions[FIRST_ROW + i].height = 20.25
    for c, w in zip("ABCDEFGH", (3, 9, 7, 18, 28, 11, 11, 20)):
        ws.column_dimensions[c].width = w
    for j in range(n):
        ws.column_dimensions[get_column_letter(FIRST_COL + j)].width = 12
    for row in ws.iter_rows(min_row=5, max_row=10, min_col=FIRST_COL, max_col=FIRST_COL + n - 1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            c.font = Font(size=8)
    ws.freeze_panes = f"{get_column_letter(FIRST_COL)}{FIRST_ROW}"


def _build_summary(ws, entries: list[GalleryEntry], sim: np.ndarray) -> None:
    """類似度シート: 縦 i > 横 j（左下半分）だけに類似度を出す（右上は左下と同じ値）。"""
    legend = [(bin_label(i), PatternFill("solid", fgColor=SIM_BINS[i][1])) for i in range(len(SIM_BINS))]

    def cell(i: int, j: int):
        if i < j:
            return None
        v = round(float(sim[i, j]), 3)
        return v, sim_fill(v), "0.000"

    _build_matrix_sheet(ws, entries, "比較結果：類似度まとめ",
                        f"画像 {len(entries)} 枚を顔 bbox のサイズ（高さ → 幅）の昇順に並べた同一人物の画像どうしの類似度（左下半分のみ）。色は類似度の段階。",
                        legend, cell)


def _build_misidentify(ws, entries: list[GalleryEntry], counts: np.ndarray, n_nontarget: int) -> None:
    """誤識別数シート: 縦 i × 横 j に「i との類似度が j との類似度より高い非対象者の画像の数」（行基準なので左下・右上とも出す）。"""
    legend = [(count_bin_label(count_bin(n)), count_fill(n)) for n in _COUNT_LEGEND]

    def cell(i: int, j: int):
        n = int(counts[i, j])
        return n, count_fill(n), "0"

    _build_matrix_sheet(ws, entries, "比較結果：誤識別画像数まとめ",
                        f"縦の画像（行）を基準に、横の画像（列）との類似度より高く出た非対象者の画像の数（非対象者 {n_nontarget} 枚の中）。"
                        "対象者どうしの類似度が高いものは数えない。色は件数の段階（0 / 1〜10 / 11〜20 / 21〜50 / 51〜100 / 以降 50 件刻み）。",
                        legend, cell)


# ---- 対象画像情報（サンプル「対象画像情報【サンプル】」）----
INFO_HEADS = ("キー情報", "カメラ\nID", "追跡\nID", "映像ファイル名", "画像ファイル名", "顔の向き", "画像サイズ\n(幅:px)", "画像サイズ\n(高さ:px)",
              "全身画像", "顔画像")
INFO_FIRST_ROW = 3
INFO_ROW_PT = 74.25
INFO_WIDTHS = (50, 9, 9, 20.625, 28, 9, 11, 11, 16, 16)  # B〜K 列
_EMU_PER_PX = 9525
_PAD_PX = 3


def _col_px(width_chars: float) -> int:
    return int(width_chars * 7 + 5)


class _FitImages:
    """セルに納まる大きさに縮小した画像（縦横比は変えない）を作る。openpyxl は保存時にファイルを読むので、保存が終わるまで消さない。"""

    def __init__(self):
        self.dir = Path(tempfile.mkdtemp(prefix="verify_face_size_xlsx_"))
        self.n = 0

    def put(self, ws, path: Path | None, col: int, row: int, cell_w: int, cell_h: int) -> bool:
        """col / row（1 始まり）のセルに path の画像を、縦横比を保ってセル内に最大化して貼る。行・列と一緒に動いて絞り込みで隠れる（twoCell）。"""
        if path is None:
            return False
        img = imread(path)
        if img is None:
            return False
        h, w = img.shape[:2]
        box_w, box_h = cell_w - 2 * _PAD_PX, cell_h - 2 * _PAD_PX
        scale = min(box_w / w, box_h / h)          # 拡大・縮小どちらもセルに納める
        dw, dh = w * scale, h * scale
        if w > 2 * box_w or h > 2 * box_h:         # 大きい画像は埋め込みを小さくする（表示の大きさはアンカーで決まる）
            k = max(2 * box_w / w, 2 * box_h / h)
            img = cv2.resize(img, (max(1, int(w * k)), max(1, int(h * k))), interpolation=cv2.INTER_AREA)
        self.n += 1
        tmp = self.dir / f"{self.n}.jpg"
        if not imwrite(tmp, img, [cv2.IMWRITE_JPEG_QUALITY, 90]):
            return False
        x0 = _PAD_PX + (box_w - dw) / 2
        y0 = _PAD_PX + (box_h - dh) / 2
        xi = XLImage(str(tmp))
        xi.anchor = TwoCellAnchor(
            editAs="twoCell",
            _from=AnchorMarker(col=col - 1, colOff=int(x0 * _EMU_PER_PX), row=row - 1, rowOff=int(y0 * _EMU_PER_PX)),
            to=AnchorMarker(col=col - 1, colOff=int((x0 + dw) * _EMU_PER_PX), row=row - 1, rowOff=int((y0 + dh) * _EMU_PER_PX)))
        ws.add_image(xi)
        return True

    def close(self) -> None:
        shutil.rmtree(self.dir, ignore_errors=True)


def _build_info_sheet(ws, entries: list[GalleryEntry], images: _FitImages) -> None:
    """検証対象人物の顔画像情報の一覧。キー情報は関数式（カメラID|追跡ID|映像ファイル名|画像ファイル名）。
    画像サイズは顔 bbox の幅・高さ（追跡IDごとの JSON から。他のシートと同じ取得元）。全身画像・顔画像はオートフィルタで隠れるよう貼る。"""
    ws.sheet_view.showGridLines = False
    for k, h in enumerate(INFO_HEADS):
        _put(ws, INFO_FIRST_ROW - 1, 2 + k, h, head=True)
    ws.row_dimensions[INFO_FIRST_ROW - 1].height = 32.25
    cell_h = int(INFO_ROW_PT * 96 / 72)
    for i, e in enumerate(entries):
        r = INFO_FIRST_ROW + i
        cam, track, video, fname, direction, w, h = _col_values(e)
        vals = [f'=C{r}&"|"&D{r}&"|"&E{r}&"|"&F{r}', cam, track, video, fname, direction, w, h, None, None]
        for k, v in enumerate(vals):
            _put(ws, r, 2 + k, v)
        ws.row_dimensions[r].height = INFO_ROW_PT
        images.put(ws, e.body_path, 10, r, _col_px(INFO_WIDTHS[8]), cell_h)   # J 列 全身画像
        images.put(ws, e.face_path, 11, r, _col_px(INFO_WIDTHS[9]), cell_h)   # K 列 顔画像
    ws.column_dimensions["A"].width = 3
    for k, w in enumerate(INFO_WIDTHS):
        ws.column_dimensions[get_column_letter(2 + k)].width = w
    last = INFO_FIRST_ROW + max(len(entries), 1) - 1
    ws.auto_filter.ref = f"B{INFO_FIRST_ROW - 1}:{get_column_letter(1 + len(INFO_HEADS))}{last}"
    ws.freeze_panes = f"B{INFO_FIRST_ROW}"


def _build_conditions(ws, cond: VerifyConditions, entries: list[GalleryEntry]) -> None:
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 80
    rows = [
        ("生成日時", cond.generated_at.strftime("%Y/%m/%d %H:%M:%S")),
        ("顔特徴量抽出モデル", cond.model_name),
        ("顔ランドマークモデル", cond.landmark_name),
        ("顔の採用基準", cond.gate_filter_text),
        ("画像数", len(entries)),
        ("追跡ID数", len({e.track_key for e in entries})),
    ]
    if cond.n_nontarget is not None:
        rows += [("比較対象フォルダ", cond.gallery_dir), ("非対象者の画像数", cond.n_nontarget),
                 ("誤認した画像の一覧（テキスト）", cond.log_text_name)]
    for i, (k, v) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")


def build_verify_face_size_report(path: Path, cond: VerifyConditions, entries: list[GalleryEntry], sim: np.ndarray | None,
                                  log: PhaseLogger, counts: np.ndarray | None = None,
                                  info_entries: list[GalleryEntry] | None = None) -> None:
    """Excel を作成して保存する。先頭に「対象画像情報」（info_entries があれば。並びは呼び出し側で決める）、sim があれば類似度シート、
    counts があれば誤識別数シートを作る。ログシートは全フェーズ終了後に追記して再保存する。"""
    images = _FitImages()
    try:
        wb = Workbook()
        wb.remove(wb.active)
        with log.phase("Excel作成・保存"):
            if info_entries is not None:
                _build_info_sheet(wb.create_sheet(INFO_SHEET_NAME), info_entries, images)
            if sim is not None:
                _build_summary(wb.create_sheet(SHEET_NAME), entries, sim)
            if counts is not None:
                _build_misidentify(wb.create_sheet(MISIDENTIFY_SHEET_NAME), entries, counts, cond.n_nontarget or 0)
            _build_conditions(wb.create_sheet("条件"), cond, entries)
            Path(path).parent.mkdir(parents=True, exist_ok=True)
            wb.save(path)
        log.finish()
        _build_log_sheet(wb.create_sheet("ログ"), log, "ログシート追記のための再保存は記録対象外")
        wb.save(path)
    finally:
        images.close()
