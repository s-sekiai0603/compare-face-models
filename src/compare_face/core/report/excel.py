"""Excel レポート生成（openpyxl）。シート: 条件 / 顔比較 / 顔比較(追跡ID集約) / 探したい人物検索結果 / ログ。

「探したい人物検索結果」は、探したい人物を指定したときだけ作る。指定した各人物が「顔比較(追跡ID集約)」の
何位に出たか（出なければその理由）を 1 行ずつ並べたもので、改修前の外見識別(reid)版と同じ形式。"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .gallery import GalleryEntry
from ..common.phase_log import PhaseLogger
from .similarity import Pair, competition_ranks

GRID_COLS = 10
COL_PX = 165
INFO_ROW_PX = 175
FACE_ROW_PX = 130
BODY_ROW_PX = 200
IMG_PAD = 4
QUERY_FILL = PatternFill("solid", fgColor="FFD9EAD3")
INVALID_FILL = PatternFill("solid", fgColor="FFBFBFBF")  # 無効画像（グレー）
TARGET_FILL = PatternFill("solid", fgColor="FFFFC0CB")  # 探したい人物（ピンク）
EXCLUDED_FILL = PatternFill("solid", fgColor="FFFCE4D6")  # クエリ自動選別で除外（薄い橙）
GATE_FILL = PatternFill("solid", fgColor="FFDDEBF7")  # 顔の採用基準（顔ゲート・品質ガード）で除外（薄い青）
MAX_TARGETS = 10
SEARCH_SHEET = "探したい人物検索結果"
SEARCH_HEADS = ["指定No", "指定(カメラID/映像/追跡ID)", "順位", "類似度", "ロケーション", "カメラID", "追跡ID",
                "映像ファイル", "撮影日時", "顔信頼度", "ヒットしたクエリ", "顔画像", "全身画像"]
SEARCH_WIDTHS = [8, 36, 8, 10, 26, 10, 10, 24, 24, 10, 22, 24, 24]

MODE_IMAGE = "顔画像指定"
MODE_TRACK = "追跡ID指定"


@dataclass
class ReportConditions:
    model_name: str      # 顔特徴量抽出モデル
    landmark_name: str   # 顔ランドマークモデル
    mode: str
    query_path: str
    gallery_dir: str
    output_dir: str
    show_body: bool      # 顔画像の下に全身画像も並べる
    generated_at: datetime
    limit: int = 500
    targets: list | None = None  # [(カメラID, 映像ファイル名, 追跡ID), ...]
    merge_sec: float = 0.0       # 追跡ID指定時の特徴量マージ単位秒（0=マージしない）
    pose_name: str = "なし"      # 頭部姿勢モデル名（HopeNet / 6DRepNet）。「なし」以外なら情報欄に yaw/pitch/roll を表示
    include_excluded: bool = False  # True: old フォルダの画像もクエリ・比較対象に含める（選別なし）
    query_filter_text: str = "指定なし"  # クエリ自動選別のしきい値（QueryFilter.describe()）
    extra_exclude_dirs: tuple = ()  # クエリ自身に加えて比較対象から外した追跡IDフォルダ（run_report の exclude_dirs）
    clahe: bool = False          # True: アライン後の顔画像に CLAHE を掛けてから特徴量化（run_report の clahe）
    gate_filter_text: str = "抽出時と同じ（既定）"  # 顔の採用基準（GateFilter.describe()）


def _norm_video(name: str) -> str:
    n = name.strip().lower()
    return n[:-4] if n.endswith(".mp4") else n


def is_target(e: GalleryEntry, targets) -> bool:
    """指定された (カメラID, 映像ファイル名, 追跡ID) のいずれかに合致するか。
    入力のある項目だけで照合する（空欄は不問）。映像ファイル名は拡張子の有無を問わない。"""
    for cam, video, track in targets or []:
        if not (cam or video or track):
            continue
        if cam and e.camera_id != cam.strip():
            continue
        if video and _norm_video(e.video_file) != _norm_video(video):
            continue
        if track and e.track_id != track.strip():
            continue
        return True
    return False


def _px_to_width(px: int) -> float:
    return max(1.0, (px - 5) / 7)


def _px_to_pt(px: int) -> float:
    return px * 0.75


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _fmt_conf(c) -> str:
    return f"{c:.3f}" if _is_num(c) else "-"


def _fmt_px(v) -> str:
    return f"{v:.0f}px" if _is_num(v) else "-"


def _fmt_ratio(v, frontal) -> str:
    if not _is_num(v):
        return "-"
    tag = "" if frontal is None else ("(正面)" if frontal else "(非正面)")
    return f"{v:.2f}{tag}"


def _fmt_quality(q) -> str:
    return f"{q:.1f}" if _is_num(q) else "-"


def _info_text(e: GalleryEntry, head: str, sim: float | None = None, extra: str = "") -> str:
    lines = [head + (f" sim={sim:.4f}" if sim is not None else "")]
    if extra:
        lines.append(extra)
    if e.invalid:
        lines += ["【無効画像】特徴量比較に未使用", f"無効理由:{e.invalid_reason or '-'}"]
    elif e.gate_excluded_reason:
        lines += ["【採用基準外】特徴量比較に未使用", f"理由:{e.gate_excluded_reason}"]
    elif e.excluded_reason:
        lines += ["【自動選別で除外】特徴量比較に未使用", f"理由:{e.excluded_reason}"]
    if not e.gate_ok and not e.gate_excluded_reason:  # 基準を緩めて使った、抽出時には不採用だった画像
        lines.append(f"抽出時は採用基準外:{e.gate_reason or '-'}")
    quality_line = f"正面度:{_fmt_ratio(e.eye_ratio, e.frontal)} / 品質:{_fmt_quality(e.quality)}"
    if e.quality_norm is not None:
        quality_line += f" / ノルム:{e.quality_norm:.1f}"
    face_line = f"顔信頼度:{_fmt_conf(e.face_conf)} / 顔幅:{_fmt_px(e.face_width)}"
    if e.crop_min_side is not None:
        face_line += f" / 切出短辺:{_fmt_px(e.crop_min_side)}"
    lines += [
        f"Loc:{e.location} / Cam:{e.camera_id}",
        f"Track:{e.track_id} / 人物信頼度:{_fmt_conf(e.person_conf)}",
        face_line,
        quality_line,
    ]
    if e.pose is not None:
        yaw, pitch, roll = (int(round(v)) for v in e.pose)
        lines.append(f"姿勢 yaw:{yaw:+d}° pitch:{pitch:+d}° roll:{roll:+d}°")
    lines += [f"映像:{e.video_file}", e.captured_text]
    return "\n".join(lines)


def _place_image(ws, cell: str, path: Path, max_w: int, max_h: int) -> None:
    try:
        img = XLImage(str(path))
    except Exception:
        return
    scale = min(max_w / img.width, max_h / img.height, 1.0)  # 縮小のみ
    img.width, img.height = max(1, int(img.width * scale)), max(1, int(img.height * scale))
    ws.add_image(img, cell)


def _setup_grid(ws) -> None:
    for c in range(1, GRID_COLS + 1):
        ws.column_dimensions[get_column_letter(c)].width = _px_to_width(COL_PX)


def _rows_per_cell(show_body: bool) -> int:
    return 3 if show_body else 2


def _write_cell(ws, start_row: int, pos: int, e: GalleryEntry, text: str,
                show_body: bool, fill: PatternFill | None = None) -> None:
    """pos 番目(0始まり)のセルを 情報行+顔画像行(+全身画像行) の組で書き込む。"""
    rows = _rows_per_cell(show_body)
    row = start_row + (pos // GRID_COLS) * rows
    col = pos % GRID_COLS + 1
    info = ws.cell(row=row, column=col, value=text)
    info.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    info.font = Font(size=9)
    cells = [info, ws.cell(row=row + 1, column=col)]
    if show_body:
        cells.append(ws.cell(row=row + 2, column=col))
    if fill:
        for c in cells:
            c.fill = fill  # 画像セルの背景も同色にする
    ws.row_dimensions[row].height = _px_to_pt(INFO_ROW_PX)
    ws.row_dimensions[row + 1].height = _px_to_pt(FACE_ROW_PX)
    letter = get_column_letter(col)
    _place_image(ws, f"{letter}{row + 1}", e.face_path, COL_PX - IMG_PAD * 2, FACE_ROW_PX - IMG_PAD * 2)
    if show_body:
        ws.row_dimensions[row + 2].height = _px_to_pt(BODY_ROW_PX)
        if e.body_path:
            _place_image(ws, f"{letter}{row + 2}", e.body_path, COL_PX - IMG_PAD * 2, BODY_ROW_PX - IMG_PAD * 2)


def _grid_rows(n: int, show_body: bool) -> int:
    return ((n + GRID_COLS - 1) // GRID_COLS) * _rows_per_cell(show_body)


def _build_conditions(ws, cond: ReportConditions, n_query: int, n_gallery: int, n_tracks: int, n_out: int,
                      n_groups: int = 0, n_invalid_query: int = 0, n_align_failed: int = 0,
                      n_landmark_missing: int | None = None, n_excluded_query: int = 0,
                      n_gate_query: int = 0, n_gate_gallery: int = 0) -> None:
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 90
    query_label = "クエリ顔画像" if cond.mode == MODE_IMAGE else "クエリ追跡IDフォルダ"
    rows = [
        ("生成日時", cond.generated_at.strftime("%Y/%m/%d %H:%M:%S")),
        ("顔特徴量抽出モデル", cond.model_name),
        ("顔ランドマークモデル", cond.landmark_name),
        ("頭部姿勢モデル（情報欄に yaw/pitch/roll を表示）", cond.pose_name),
        ("クエリ指定方法", cond.mode),
        (query_label, cond.query_path),
        ("比較対象フォルダ", cond.gallery_dir),
        ("レポート出力先フォルダ", cond.output_dir),
        ("全身画像の表示", "ON" if cond.show_body else "OFF"),
        ("レポート出力件数(上位)", cond.limit),
        ("old フォルダの画像", "含める（選別なし）" if cond.include_excluded else "除外（選別あり。既定）"),
        ("顔の採用基準（抽出時の顔ゲート・品質ガード。クエリ・比較対象に適用）", cond.gate_filter_text),
        ("比較対象から追加で除外した追跡IDフォルダ", "\n".join(cond.extra_exclude_dirs) or "指定なし"),
        ("探したい人物(ピンク表示)", "\n".join(
            f"カメラID:{c or '-'} / 映像:{v or '-'} / 追跡ID:{t or '-'}"
            for c, v, t in (cond.targets or []) if c or v or t) or "指定なし"),
        ("クエリ画像数", n_query),
        ("クエリ無効画像数（比較に不使用）", n_invalid_query),
        ("クエリ採用基準外の画像数（比較に不使用）", n_gate_query),
        ("クエリ自動選別（しきい値）", cond.query_filter_text),
        ("前処理 CLAHE（明るさ補正）", "あり（アライン後の顔画像に適用）" if cond.clahe else "なし"),
        ("クエリ自動選別で除外した画像数（比較に不使用）", n_excluded_query),
        ("特徴量マージ単位(秒)", (f"{cond.merge_sec:g}（クエリグループ数: {n_groups}）"
                              if cond.merge_sec > 0 else "指定なし")),
        ("比較対象 画像数", n_gallery),
        ("比較対象 採用基準外の画像数（除外）", n_gate_gallery),
        ("比較対象 追跡ID数", n_tracks),
        ("アライン失敗数（リサイズで代替）", n_align_failed),
        ("ランドマーク未取得数（テンプレート座標で代替）",
         "対象外（ランドマーク不要のモデル）" if n_landmark_missing is None else n_landmark_missing),
        ("顔比較シート 出力件数", n_out),
    ]
    for i, (k, v) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)


def _build_compare_sheet(ws, mode: str, queries: list[GalleryEntry], gallery: list[GalleryEntry],
                         pairs: list[Pair], show_body: bool, targets=None,
                         group_ids: list[int | None] | None = None,
                         used_queries: list[GalleryEntry] | None = None) -> None:
    """queries=表示する全クエリ画像(無効含む)。pair.q_idx は used_queries(無効を除く)または
    グループ(ID-1)を指す。group_ids は queries と同順で、無効画像は None。"""
    used_queries = queries if used_queries is None else used_queries
    _setup_grid(ws)
    for i, q in enumerate(queries):
        gid = group_ids[i] if group_ids else None
        gtext = f"クエリグループID:G{gid}" if gid else ""
        fill = (INVALID_FILL if q.invalid else GATE_FILL if q.gate_excluded_reason
                else EXCLUDED_FILL if q.excluded_reason else QUERY_FILL)
        _write_cell(ws, 1, i, q, _info_text(q, "クエリ", extra=gtext), show_body, fill=fill)
    start = 1 + _grid_rows(len(queries), show_body) + 1  # 1行空ける
    ranks = competition_ranks([p.sim for p in pairs])
    for i, (p, rank) in enumerate(zip(pairs, ranks)):
        g = gallery[p.g_idx]
        if group_ids is not None:  # マージ時は pair.q_idx がグループ(ID-1)を指す
            extra = f"ヒットしたクエリグループID:G{p.q_idx + 1}"
        else:
            extra = f"Query:{used_queries[p.q_idx].captured_text}" if mode == MODE_TRACK else ""
        _write_cell(ws, start, i, g, _info_text(g, f"{rank}位", p.sim, extra), show_body,
                    fill=TARGET_FILL if is_target(g, targets) else None)


SELF_QUERY_MESSAGE = "クエリ自身のため比較対象から除外（照合対象外）"
EXCLUDED_DIR_MESSAGE = "比較対象から追加で除外した追跡IDフォルダのため照合対象外（クエリ画像の元の追跡ID）"


def _dir_matches_spec(dir_path: str, spec) -> bool:
    """追跡IDフォルダのパス（…/カメラID/映像名/追跡ID）が探したい人物の指定に合致するか（空欄は不問）。"""
    parts = Path(dir_path).parts
    if len(parts) < 3:
        return False
    cam, video, track = spec
    if cam and parts[-3] != cam.strip():
        return False
    if video and _norm_video(parts[-2]) != _norm_video(video):
        return False
    if track and parts[-1] != track.strip():
        return False
    return bool(cam or video or track)


def search_results(targets, gallery: list[GalleryEntry], aggregated: list[Pair], limit: int,
                   group_ids: list[int | None] | None = None, used_queries: list[GalleryEntry] | None = None,
                   mode: str = MODE_TRACK, queries: list[GalleryEntry] | None = None,
                   excluded_dirs=()) -> list[dict]:
    """探したい人物ごとに、追跡ID集約の上位 limit 件での最上位ヒットを求める。
    戻り値の各要素: {"no", "spec", "rank", "sim", "entry", "hit", "message"}（ヒットなしは rank/sim/entry が None）。
    指定がクエリ自身（比較対象から除外されている）に当たる場合は SELF_QUERY_MESSAGE、
    excluded_dirs（追加で除外した追跡IDフォルダ）に当たる場合は EXCLUDED_DIR_MESSAGE を入れる。"""
    ranks = competition_ranks([p.sim for p in aggregated])
    out = []
    for no, spec in enumerate((t for t in (targets or []) if any(t)), start=1):
        cam, video, track = spec
        row = {"no": no, "spec": f"{cam or '-'} / {video or '-'} / {track or '-'}",
               "rank": None, "sim": None, "entry": None, "hit": "", "message": ""}
        for p, rank in zip(aggregated, ranks):
            g = gallery[p.g_idx]
            if not is_target(g, [spec]):
                continue
            row.update(rank=rank, sim=p.sim, entry=g)
            if group_ids is not None:
                row["hit"] = f"G{p.q_idx + 1}"
            elif mode == MODE_TRACK and used_queries:
                row["hit"] = used_queries[p.q_idx].captured_text
            else:
                row["hit"] = "-"
            break
        if row["entry"] is None:
            if any(is_target(g, [spec]) for g in gallery):
                row["message"] = f"比較対象には存在するが、追跡ID集約の上位{limit}件に入っていません"
            elif any(is_target(q, [spec]) for q in (queries or [])):
                row["message"] = SELF_QUERY_MESSAGE
            elif any(_dir_matches_spec(d, spec) for d in excluded_dirs):
                row["message"] = EXCLUDED_DIR_MESSAGE
            else:
                row["message"] = "比較対象フォルダに該当する人物がありません（無効画像・採用基準外のみの場合を含む）"
        out.append(row)
    return out


def _build_search_sheet(ws, results: list[dict], show_body: bool) -> None:
    heads = SEARCH_HEADS if show_body else SEARCH_HEADS[:-1]
    for c, (h, w) in enumerate(zip(heads, SEARCH_WIDTHS), start=1):
        ws.cell(row=1, column=c, value=h).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    for i, r in enumerate(results, start=2):
        ws.cell(row=i, column=1, value=r["no"])
        ws.cell(row=i, column=2, value=r["spec"])
        e = r["entry"]
        if e is None:
            ws.cell(row=i, column=3, value=r["message"])
            continue
        values = [r["rank"], round(float(r["sim"]), 4), e.location, e.camera_id, e.track_id, e.video_file,
                  e.captured_text, _fmt_conf(e.face_conf), r["hit"]]
        for c, v in enumerate(values, start=3):
            cell = ws.cell(row=i, column=c, value=v)
            cell.fill = TARGET_FILL
        row_px = BODY_ROW_PX if show_body else FACE_ROW_PX
        ws.row_dimensions[i].height = _px_to_pt(row_px)
        img_w = int(SEARCH_WIDTHS[11] * 7 + 5) - IMG_PAD * 2
        _place_image(ws, f"{get_column_letter(12)}{i}", e.face_path, img_w, row_px - IMG_PAD * 2)
        if show_body and e.body_path:
            _place_image(ws, f"{get_column_letter(13)}{i}", e.body_path, img_w, row_px - IMG_PAD * 2)


def _build_log_sheet(ws, log: PhaseLogger, note: str = "") -> None:
    heads = ["No", "フェーズ", "開始時刻", "終了時刻", "所要時間(秒)", "備考"]
    widths = [6, 50, 24, 24, 14, 70]
    for c, (h, w) in enumerate(zip(heads, widths), start=1):
        ws.cell(row=1, column=c, value=h).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(c)].width = w

    def ms(dt: datetime) -> str:
        return dt.strftime("%Y/%m/%d %H:%M:%S.") + f"{dt.microsecond // 1000:03d}"

    row = 2
    ws.cell(row=row, column=2, value="レポート出力 全体")
    ws.cell(row=row, column=3, value=ms(log.started_at))
    if log.finished_at:
        ws.cell(row=row, column=4, value=ms(log.finished_at))
        ws.cell(row=row, column=5, value=round((log.finished_at - log.started_at).total_seconds(), 3))
    ws.cell(row=row, column=6, value=note)
    for rec in log.records:
        row += 1
        ws.cell(row=row, column=1, value=row - 2)
        ws.cell(row=row, column=2, value=rec.name)
        ws.cell(row=row, column=3, value=ms(rec.start))
        if rec.end:
            ws.cell(row=row, column=4, value=ms(rec.end))
            ws.cell(row=row, column=5, value=round(rec.elapsed_sec, 3))
        ws.cell(row=row, column=6, value=rec.note)


def build_report(path: Path, cond: ReportConditions, queries: list[GalleryEntry],
                 gallery: list[GalleryEntry], ranked: list[Pair], aggregated: list[Pair],
                 log: PhaseLogger, group_ids: list[int | None] | None = None,
                 used_queries: list[GalleryEntry] | None = None, n_align_failed: int = 0,
                 n_landmark_missing: int | None = None, n_gate_gallery: int = 0) -> None:
    """レポートを作成して保存する。ログシートは全フェーズ終了後に追記して再保存する。
    n_gate_gallery は顔の採用基準で比較対象から外した画像数（gallery には含まれないので別に渡す）。"""
    wb = Workbook()
    ws_cond = wb.active
    ws_cond.title = "条件"
    n_tracks = len({g.track_key for g in gallery})
    _build_conditions(ws_cond, cond, len(queries), len(gallery), n_tracks, len(ranked),
                      max((g for g in (group_ids or []) if g), default=0),
                      sum(1 for q in queries if q.invalid), n_align_failed, n_landmark_missing,
                      sum(1 for q in queries if q.excluded_reason and not q.invalid and not q.gate_excluded_reason),
                      sum(1 for q in queries if q.gate_excluded_reason and not q.invalid), n_gate_gallery)
    with log.phase("Excel作成（シート生成・画像配置）"):
        _build_compare_sheet(wb.create_sheet("顔比較"), cond.mode, queries, gallery, ranked,
                             cond.show_body, cond.targets, group_ids, used_queries)
        _build_compare_sheet(wb.create_sheet("顔比較(追跡ID集約)"), cond.mode, queries, gallery,
                             aggregated, cond.show_body, cond.targets, group_ids, used_queries)
        if any(any(t) for t in (cond.targets or [])):
            _build_search_sheet(wb.create_sheet(SEARCH_SHEET),
                                search_results(cond.targets, gallery, aggregated, cond.limit, group_ids,
                                               used_queries, cond.mode, queries, cond.extra_exclude_dirs),
                                cond.show_body)
    with log.phase("Excel保存"):
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
    log.finish()
    _build_log_sheet(wb.create_sheet("ログ"), log, "ログシート追記のための再保存は記録対象外")
    wb.save(path)
