"""検索元セット検索の Excel 出力。
シート: 条件 / 検索元 / 総当たり（表） / 総当たり（一覧） / 全体検索（画像） / 全体検索（追跡ID） / ログ。
画像は高さ THUMB_H px の縮小版を一時フォルダに作って貼る（総当たりは N×(N-1) 行 × 2 枚になるため）。"""
from __future__ import annotations

import shutil
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import cv2
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..common.imageio import imread, imwrite
from ..common.phase_log import PhaseLogger
from .excel import _build_log_sheet, _fmt_conf, _fmt_px, _fmt_ratio
from .gallery import GalleryEntry
from .query_set import QueryImage, SearchResult

THUMB_H = 56
ROW_PX = 62
HEAD_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFFFF")
OWN_FILL = PatternFill("solid", fgColor="FFE7E6E6")   # 検索元自身の追跡ID（集計から除く）
RANK_FILLS = ((10, PatternFill("solid", fgColor="FFC6EFCE")),     # 10 位以内: 緑
              (100, PatternFill("solid", fgColor="FFFFEB9C")),    # 100 位以内: 黄
              (None, PatternFill("solid", fgColor="FFF8CBAD")))   # それ以上: 橙
_THIN = Side(style="thin", color="FFBFBFBF")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
CENTER = Alignment(wrap_text=True, horizontal="center", vertical="center")
NA = "-"


@dataclass
class QuerySetConditions:
    model_name: str
    landmark_name: str
    pose_name: str
    set_dir: str
    gallery_dir: str
    output_dir: str
    generated_at: datetime
    include_excluded: bool = False
    gate_filter_text: str = "抽出時と同じ（既定）"
    own_specs: list | None = None   # 本人の追跡IDの追加指定 [(カメラID, 映像ファイル名, 追跡ID), ...]
    use_cache: bool = True


def rank_fill(rank) -> PatternFill | None:
    if not isinstance(rank, int):
        return None
    for limit, fill in RANK_FILLS:
        if limit is None or rank <= limit:
            return fill
    return None


class _Thumbs:
    """貼り付け用の縮小画像（高さ THUMB_H）。openpyxl は保存時にファイルを読むので、保存が終わるまで消さない。"""

    def __init__(self):
        self.dir = Path(tempfile.mkdtemp(prefix="query_set_xlsx_"))
        self._cache: dict[Path, tuple[Path, int, int] | None] = {}

    def image(self, path: Path) -> XLImage | None:
        if path not in self._cache:
            img = imread(path)
            if img is None:
                self._cache[path] = None
            else:
                h, w = img.shape[:2]
                scale = min(1.0, THUMB_H / max(1, h))
                nw, nh = max(1, int(w * scale)), max(1, int(h * scale))
                if scale < 1.0:
                    img = cv2.resize(img, (nw, nh), interpolation=cv2.INTER_AREA)
                p = self.dir / f"{len(self._cache)}.jpg"
                self._cache[path] = (p, nw, nh) if imwrite(p, img, [cv2.IMWRITE_JPEG_QUALITY, 90]) else None  # 書けなければ画像なし
        v = self._cache[path]
        if v is None:
            return None
        xi = XLImage(str(v[0]))
        xi.width, xi.height = v[1], v[2]
        return xi

    def close(self) -> None:
        shutil.rmtree(self.dir, ignore_errors=True)


def _put(ws, r: int, c: int, v, fill: PatternFill | None = None, head: bool = False, bold: bool = False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.border = BORDER
    cell.alignment = CENTER
    if head:
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    else:
        if fill is not None:
            cell.fill = fill
        if bold:
            cell.font = Font(bold=True)
    return cell


def _heads(ws, r: int, heads: list[str], height: float = 45) -> None:
    for c, h in enumerate(heads, start=1):
        _put(ws, r, c, h, head=True)
    ws.row_dimensions[r].height = height


def _widths(ws, widths: list[float]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _pose_text(e: GalleryEntry) -> str:
    if e.pose is None:
        return NA
    yaw, pitch, roll = (int(round(v)) for v in e.pose)
    return f"yaw:{yaw:+d}° pitch:{pitch:+d}° roll:{roll:+d}°"


def _origin_text(q: QueryImage) -> str:
    e = q.entry
    return f"{e.camera_id} / {e.video_file} / {e.track_id}" if q.original is not None else "-（元画像なし）"


def _query_text(q: QueryImage) -> str:
    return f"No{q.no}\n{q.label}\n{_origin_text(q)}"


def _image_at(ws, thumbs: _Thumbs, r: int, c: int, path: Path) -> None:
    xi = thumbs.image(path)
    if xi is not None:
        ws.add_image(xi, f"{get_column_letter(c)}{r}")


def _query_image_path(q: QueryImage) -> Path:
    return q.path


# ---- 各シート ----
def _build_conditions(ws, cond: QuerySetConditions, queries: list[QueryImage], gallery: list[GalleryEntry],
                      result: SearchResult, n_gate_gallery: int) -> None:
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90
    own = "\n".join(f"{c} / {v} / {t}" for _, c, v, t in result.own_tracks) or "なし"
    specs = "\n".join(f"カメラID:{c or '-'} / 映像:{v or '-'} / 追跡ID:{t or '-'}" for c, v, t in (cond.own_specs or [])) or "指定なし"
    unmatched = "\n".join(f"カメラID:{c or '-'} / 映像:{v or '-'} / 追跡ID:{t or '-'}" for c, v, t in result.unmatched_specs)
    rows = [
        ("生成日時", cond.generated_at.strftime("%Y/%m/%d %H:%M:%S")),
        ("顔特徴量抽出モデル", cond.model_name),
        ("顔ランドマークモデル", cond.landmark_name),
        ("頭部姿勢モデル（検索元の yaw/pitch/roll）", cond.pose_name),
        ("検索元セット", cond.set_dir),
        ("比較対象フォルダ", cond.gallery_dir),
        ("レポート出力先フォルダ", cond.output_dir),
        ("old フォルダの画像", "含める（選別なし）" if cond.include_excluded else "除外（選別あり。既定）"),
        ("顔の採用基準（抽出時の顔ゲート・品質ガード。比較対象に適用）", cond.gate_filter_text),
        ("特徴量キャッシュ", "使う" if cond.use_cache else "使わない"),
        ("検索元 画像数", len(queries)),
        ("検索元 グループ数", len({q.group for q in queries})),
        ("検索元 元画像を特定した画像数", sum(1 for q in queries if q.original is not None)),
        ("検索元 比較対象に含まれる画像数", sum(1 for q in queries if q.gallery_index is not None)),
        ("本人の追跡ID（検索元の元画像の追跡ID + 追加指定）", own),
        ("本人の追跡IDの追加指定", specs),
        ("追加指定のうち比較対象に無かったもの", unmatched or "なし"),
        ("比較対象 画像数", len(gallery)),
        ("比較対象 採用基準外の画像数（除外）", n_gate_gallery),
        ("比較対象 追跡ID数", result.n_tracks),
        ("総当たりの組数", len(queries) * max(0, len(queries) - 1)),
    ]
    for i, (k, v) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")


def _build_queries_sheet(ws, queries: list[QueryImage], thumbs: _Thumbs) -> None:
    ws.cell(row=1, column=1, value="検索元セットの画像。元画像 = 比較対象フォルダ内の同じ画像（バイト一致）。カメラID・追跡ID・顔信頼度などは元画像の記録から。"
                                  "「比較対象での扱い」が「含む」以外の画像は、全体検索でこの画像を検索対象にした順位は出ない（検索元としては使う）。").font = Font(size=9)
    heads = ["No", "グループ", "ファイル名", "画像", "カメラID", "映像ファイル", "追跡ID", "撮影日時", "顔信頼度", "顔幅",
             "正面度", "頭部姿勢", "抽出時の採用基準", "比較対象での扱い"]
    _heads(ws, 3, heads)
    for i, q in enumerate(queries):
        r = 4 + i
        e = q.entry
        matched = q.original is not None
        vals = [q.no, q.group or NA, q.path.name, "", e.camera_id if matched else NA, e.video_file if matched else NA,
                e.track_id if matched else NA, e.captured_text if matched else NA, _fmt_conf(e.face_conf), _fmt_px(e.face_width),
                _fmt_ratio(e.eye_ratio, e.frontal), _pose_text(e),
                ("採用" if e.gate_ok else f"採用基準外: {e.gate_reason or NA}") if matched else NA, q.treatment or NA]
        for c, v in enumerate(vals, start=1):
            _put(ws, r, c, v)
        _image_at(ws, thumbs, r, 4, _query_image_path(q))
        ws.row_dimensions[r].height = ROW_PX * 0.75
    _widths(ws, [6, 16, 24, 10, 10, 22, 9, 22, 10, 9, 14, 26, 30, 44])
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(heads))}{3 + len(queries)}"


def _build_matrix_sheet(ws, queries: list[QueryImage], result: SearchResult, thumbs: _Thumbs) -> None:
    n = len(queries)
    ws.cell(row=1, column=1, value=f"検索元セット {n} 枚の総当たりの類似度（行 = 検索元、列 = 検索対象。対角は自分自身）。"
                                  "色は全体の中での相対値（赤 低 〜 緑 高）。").font = Font(size=9)
    _put(ws, 3, 1, "No", head=True)
    _put(ws, 3, 2, "検索元", head=True)
    _put(ws, 3, 3, "画像", head=True)
    for j, q in enumerate(queries):
        _put(ws, 3, 4 + j, f"No{q.no}\n{q.label}", head=True)
    ws.row_dimensions[3].height = 60
    for i, q in enumerate(queries):
        r = 4 + i
        _put(ws, r, 1, q.no, bold=True)
        _put(ws, r, 2, f"{q.label}\n{_origin_text(q)}")
        _image_at(ws, thumbs, r, 3, _query_image_path(q))
        for j in range(n):
            _put(ws, r, 4 + j, NA if i == j else round(float(result.round_robin[i, j]), 3))
        ws.row_dimensions[r].height = ROW_PX * 0.75
    if n > 1:
        rng = f"D4:{get_column_letter(3 + n)}{3 + n}"
        ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="FFF8696B",
                                                          mid_type="percentile", mid_value=50, mid_color="FFFFEB84",
                                                          end_type="max", end_color="FF63BE7B"))
    _widths(ws, [6, 24, 10] + [9] * n)
    ws.freeze_panes = "D4"


def _build_pairs_sheet(ws, queries: list[QueryImage], result: SearchResult, thumbs: _Thumbs) -> None:
    ws.cell(row=1, column=1, value="総当たり（一覧）: 各画像を検索元にして、セット内の自分以外の全画像と 1 枚ずつ比べた類似度。"
                                  "同じグループ = 同じサブフォルダ（同じ追跡IDなど、人が同じ条件と決めた組）。").font = Font(size=9)
    heads = ["組 No", "検索元 No", "検索元", "検索元画像", "検索元の頭部姿勢", "検索対象 No", "検索対象", "検索対象画像",
             "検索対象の頭部姿勢", "類似度", "同じグループ"]
    _heads(ws, 3, heads)
    r, no = 4, 0
    for i, q in enumerate(queries):
        for j, g in enumerate(queries):
            if i == j:
                continue
            no += 1
            vals = [no, q.no, f"{q.label}\n{_origin_text(q)}", "", _pose_text(q.entry), g.no,
                    f"{g.label}\n{_origin_text(g)}", "", _pose_text(g.entry), round(float(result.round_robin[i, j]), 3),
                    "○" if q.group and q.group == g.group else ""]
            for c, v in enumerate(vals, start=1):
                _put(ws, r, c, v)
            _image_at(ws, thumbs, r, 4, _query_image_path(q))
            _image_at(ws, thumbs, r, 8, _query_image_path(g))
            ws.row_dimensions[r].height = ROW_PX * 0.75
            r += 1
    _widths(ws, [7, 9, 26, 10, 24, 9, 26, 10, 24, 9, 11])
    ws.freeze_panes = "A4"
    if r > 4:
        ws.auto_filter.ref = f"A3:{get_column_letter(len(heads))}{r - 1}"


def _build_image_rank_sheet(ws, queries: list[QueryImage], result: SearchResult, thumbs: _Thumbs) -> None:
    ws.cell(row=1, column=1, value="全体検索（画像）: 各画像を検索元にして比較対象フォルダ全体で検索したとき、セット内の他の画像（検索対象）とその追跡IDが何位に出たか。"
                                  "検索元自身の画像 1 枚だけを比較対象から外す（レポートと同じ）。").font = Font(size=9)
    ws.cell(row=2, column=1, value="画像の順位 = 比較対象の全画像の中での順位（レポートの「顔比較」）。追跡IDの順位 = 追跡IDごとに一番似た画像で並べたときの順位"
                                  "（「顔比較(追跡ID集約)」）。カメラ別 = 検索対象と同じカメラの中での順位。他人だけ = 本人の追跡IDを数えない順位。"
                                  "色: 緑 10 位以内 / 黄 100 位以内 / 橙 それ以上。").font = Font(size=9)
    heads = ["組 No", "検索元 No", "検索元", "検索元画像", "検索対象 No", "検索対象", "検索対象画像", "類似度",
             "画像の順位\n（カメラ別）", "画像の順位\n（全体）", "追跡IDの順位\n（カメラ別）", "追跡IDの順位\n（全体）", "追跡IDの順位\n（他人だけ）"]
    _heads(ws, 4, heads)
    r, no = 5, 0
    for i, q in enumerate(queries):
        for j, g in enumerate(queries):
            if i == j:
                continue
            no += 1
            ir = result.image_ranks.get((i, j))
            tr = result.track_ranks.get((i, g.entry.track_key)) if g.original is not None else None
            sim = round(float(result.round_robin[i, j]), 3)
            same_image = q.gallery_index is not None and g.gallery_index == q.gallery_index  # 同じ元画像を 2 か所に置いた
            no_rank = "検索元と同じ画像" if same_image else "比較対象に無い"
            vals = [no, q.no, f"{q.label}\n{_origin_text(q)}", "", g.no, f"{g.label}\n{_origin_text(g)}", "", sim,
                    ir.rank_camera if ir else no_rank, ir.rank_all if ir else NA,
                    tr.rank_camera if tr else NA, tr.rank_all if tr else NA, tr.rank_strangers if tr else NA]
            for c, v in enumerate(vals, start=1):
                _put(ws, r, c, v, fill=rank_fill(v) if c >= 9 else None)
            _image_at(ws, thumbs, r, 4, _query_image_path(q))
            _image_at(ws, thumbs, r, 7, _query_image_path(g))
            ws.row_dimensions[r].height = ROW_PX * 0.75
            r += 1
    _widths(ws, [7, 9, 26, 10, 9, 26, 10, 9, 12, 12, 12, 12, 12])
    ws.freeze_panes = "A5"
    if r > 5:
        ws.auto_filter.ref = f"A4:{get_column_letter(len(heads))}{r - 1}"


def _build_track_rank_sheet(ws, queries: list[QueryImage], result: SearchResult, thumbs: _Thumbs) -> None:
    own = result.own_tracks
    ws.cell(row=1, column=1, value="全体検索（追跡ID）: 検索元の画像ごとに、本人の各追跡IDが比較対象全体で何位に出たか（追跡ID単位）。"
                                  "本人の追跡ID = 検索元の元画像の追跡ID + 追加指定。灰色 = 検索元自身の追跡ID（下の集計から除く）。").font = Font(size=9)
    ws.cell(row=2, column=1, value="類似度 = その追跡IDで一番似た画像の類似度。カメラ別 = 同じカメラの追跡IDの中での順位。全体 = 全追跡IDの中での順位。"
                                  "他人だけ = 本人の追跡IDを数えずに、他人の追跡IDの中で何位か。空欄「-」= その追跡IDが比較対象に無い（検索元自身の画像しか無い場合を含む）。").font = Font(size=9)
    _put(ws, 4, 1, "No", head=True)
    _put(ws, 4, 2, "検索元", head=True)
    _put(ws, 4, 3, "画像", head=True)
    for c in range(1, 4):
        ws.merge_cells(start_row=4, start_column=c, end_row=5, end_column=c)
        _put(ws, 5, c, None, head=True)
    sub = ("類似度", "カメラ別", "全体", "他人だけ")
    for k, (_, cam, video, tid) in enumerate(own):
        c0 = 4 + k * len(sub)
        _put(ws, 4, c0, f"{cam} {tid}\n{video}", head=True)
        for d in range(1, len(sub)):
            _put(ws, 4, c0 + d, None, head=True)
        ws.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + len(sub) - 1)
        for d, t in enumerate(sub):
            _put(ws, 5, c0 + d, t, head=True)
    ws.row_dimensions[4].height = 40
    r = 6
    for i, q in enumerate(queries):
        _put(ws, r, 1, q.no, bold=True)
        _put(ws, r, 2, f"{q.label}\n{_origin_text(q)}")
        _image_at(ws, thumbs, r, 3, _query_image_path(q))
        for k, (key, _, _, _) in enumerate(own):
            c0 = 4 + k * len(sub)
            tr = result.track_ranks.get((i, key))
            is_own = q.original is not None and q.entry.track_key == key
            vals = (round(tr.sim, 3), tr.rank_camera, tr.rank_all, tr.rank_strangers) if tr else (NA, NA, NA, NA)
            for d, v in enumerate(vals):
                _put(ws, r, c0 + d, v, fill=OWN_FILL if is_own else (rank_fill(v) if d >= 1 else None))
        ws.row_dimensions[r].height = ROW_PX * 0.75
        r += 1
    # 集計（検索元自身の追跡IDは除く）
    r += 1
    _put(ws, r, 1, "集計", head=True)
    _put(ws, r, 2, "他人だけの順位（検索元自身の追跡IDを除く）", head=True)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    _put(ws, r, 3, None, head=True)
    for k in range(len(own)):
        c0 = 4 + k * len(sub)
        for d in range(len(sub)):
            _put(ws, r, c0 + d, None, head=True)
        ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + len(sub) - 1)
        ws.cell(row=r, column=c0, value=f"{own[k][1]} {own[k][3]}")
    for label, limit in (("1 位", 1), ("10 位以内", 10), ("100 位以内", 100)):
        r += 1
        _put(ws, r, 1, None)
        _put(ws, r, 2, label, bold=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        _put(ws, r, 3, None)
        for k, (key, _, _, _) in enumerate(own):
            c0 = 4 + k * len(sub)
            vals = [result.track_ranks[(i, key)].rank_strangers for i, q in enumerate(queries)
                    if (i, key) in result.track_ranks and not (q.original is not None and q.entry.track_key == key)]
            text = f"{sum(1 for v in vals if v <= limit)} / {len(vals)}" if vals else NA
            _put(ws, r, c0, text)
            for d in range(1, len(sub)):
                _put(ws, r, c0 + d, None)
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + len(sub) - 1)
    _widths(ws, [6, 26, 10] + [9] * (len(sub) * len(own)))
    ws.freeze_panes = "D6"


def build_query_set_report(path: Path, cond: QuerySetConditions, queries: list[QueryImage], gallery: list[GalleryEntry],
                           result: SearchResult, log: PhaseLogger, n_gate_gallery: int = 0) -> None:
    """Excel を作成して保存する。ログシートは全フェーズ終了後に追記して再保存する。"""
    thumbs = _Thumbs()
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "条件"
        _build_conditions(ws, cond, queries, gallery, result, n_gate_gallery)
        with log.phase("Excel作成（シート生成・画像配置）"):
            _build_queries_sheet(wb.create_sheet("検索元"), queries, thumbs)
            _build_matrix_sheet(wb.create_sheet("総当たり（表）"), queries, result, thumbs)
            _build_pairs_sheet(wb.create_sheet("総当たり（一覧）"), queries, result, thumbs)
            _build_image_rank_sheet(wb.create_sheet("全体検索（画像）"), queries, result, thumbs)
            _build_track_rank_sheet(wb.create_sheet("全体検索（追跡ID）"), queries, result, thumbs)
        with log.phase("Excel保存"):
            Path(path).parent.mkdir(parents=True, exist_ok=True)
            wb.save(path)
        log.finish()
        _build_log_sheet(wb.create_sheet("ログ"), log, "ログシート追記のための再保存は記録対象外")
        wb.save(path)
    finally:
        thumbs.close()
