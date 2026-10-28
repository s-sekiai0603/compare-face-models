"""検索元セット検索（core/report/query_set.py, query_set_excel.py）。GPU・モデル重みなしで通る。"""
import json
import shutil
import zlib
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
import openpyxl
import pytest

from compare_face.core.common.imageio import imwrite  # 日本語パス対応（cv2.imwrite は非 ASCII パスに書けない）
from compare_face.core.common.phase_log import PhaseLogger
from compare_face.core.report.gallery import QUERY_SETS_DIR_NAME, GalleryEntry, scan_gallery
from compare_face.core.report.query_set import (
    QueryImage, compute_search, create_query_set, default_query_sets_root, list_query_sets, match_originals,
    parse_track_lines, run_query_set_search, scan_query_images,
)
from compare_face.core.report.query_set_excel import QuerySetConditions, build_query_set_report, rank_fill


def _img(seed: int, size=(40, 36)) -> np.ndarray:
    rng = np.random.default_rng(seed)
    return rng.integers(0, 255, (size[0], size[1], 3), np.uint8)


def _make_track(root, run, cam, video, tid, stamps, **fields):
    """追跡IDフォルダを作る。画像は stamp ごとに内容が違う（バイト一致で元画像を特定するため。seed は決定的に 10000〜）。"""
    d = root / run / cam / video / str(tid)
    d.mkdir(parents=True, exist_ok=True)
    images = []
    for ts in stamps:
        assert imwrite(d / f"{ts}_face.jpg", _img(10000 + zlib.crc32(f"{cam}/{tid}/{ts}".encode()) % 100000))
        images.append({"file_face": f"{ts}_face.jpg", "file_body": None, "person_conf": 0.9, "face_conf": 0.8,
                       "face_width": 42.0, "eye_ratio": 0.31, "frontal": True, "quality": 5.0,
                       "face_bbox_xyxy": [10, 5, 52, 50], "time_sec": 1.0, **fields})
    (d / f"{tid}.json").write_text(json.dumps(
        {"track_id": tid, "video_file": f"{video}.mp4", "location": "L", "images": images}, ensure_ascii=False), encoding="utf-8")
    return d


# ---- セットの置き場所・走査 ----
def test_create_list_and_scan_sets(tmp_path):
    root = default_query_sets_root(tmp_path / "faces")
    assert root == tmp_path / "faces" / QUERY_SETS_DIR_NAME
    assert list_query_sets(root) == []
    d = create_query_set(root, " Aさん ")
    assert d == root / "Aさん" and d.is_dir()
    with pytest.raises(ValueError):
        create_query_set(root, "a/b")
    with pytest.raises(ValueError):
        create_query_set(root, "  ")
    (d / "Cam4_79").mkdir()
    (d / "Cam1_463" / "深い").mkdir(parents=True)
    assert imwrite(d / "Cam4_79" / "02_右横顔.jpg", _img(1))
    assert imwrite(d / "Cam4_79" / "01_正面.JPG", _img(2))
    assert imwrite(d / "Cam1_463" / "深い" / "x.png", _img(3))
    assert imwrite(d / "直下.jpg", _img(4))
    (d / "memo.txt").write_text("x", encoding="utf-8")
    found = scan_query_images(d)
    assert [(g, p.name) for g, p in found] == [("", "直下.jpg"), ("Cam1_463", "x.png"), ("Cam4_79", "01_正面.JPG"), ("Cam4_79", "02_右横顔.jpg")]
    create_query_set(root, "空")
    assert list_query_sets(root) == [("Aさん", 4), ("空", 0)]
    assert scan_query_images(tmp_path / "ない") == []


def test_scan_gallery_skips_query_sets_dir(tmp_path):
    d = _make_track(tmp_path, "r", "Cam1", "v", 1, ["20260311_122723_000"])
    s = tmp_path / QUERY_SETS_DIR_NAME / "S" / "g"
    s.mkdir(parents=True)
    shutil.copy2(d / "20260311_122723_000_face.jpg", s / "20260311_122723_000_face.jpg")
    assert len(scan_gallery(tmp_path)) == 1 and len(scan_gallery(tmp_path, include_excluded=True)) == 1
    assert [e.ts for e in scan_gallery(tmp_path)] == ["20260311_122723_000"] and scan_gallery(tmp_path)[0].track_id == "1"
    # 走査ルート直下だけを飛ばす。深い階層の同名フォルダ（カメラ名・追跡IDが「検索元」）は対象のまま
    _make_track(tmp_path, "r", QUERY_SETS_DIR_NAME, "v", 9, ["20260311_122730_000"])
    assert sorted(e.track_id for e in scan_gallery(tmp_path)) == ["1", "9"]


def test_match_originals_by_bytes(tmp_path):
    d = _make_track(tmp_path, "r", "Cam1", "v", 1, ["20260311_122723_000", "20260311_122724_000"])
    s = tmp_path / QUERY_SETS_DIR_NAME / "S"
    s.mkdir(parents=True)
    shutil.copy2(d / "20260311_122724_000_face.jpg", s / "名前を変えた.jpg")   # 名前を変えても内容で見つかる
    assert imwrite(s / "別の画像.jpg", _img(999))                                # 比較対象に無い画像
    entries = scan_gallery(tmp_path)
    got = match_originals([s / "名前を変えた.jpg", s / "別の画像.jpg", s / "存在しない.jpg"], entries)
    assert list(got) == [s / "名前を変えた.jpg"] and got[s / "名前を変えた.jpg"].ts == "20260311_122724_000"


def test_parse_track_lines():
    text = "Cam1|20260311113522.mp4|463\n\n# コメント\n Cam2 || 172 \n|||\n5"
    assert parse_track_lines(text) == [("Cam1", "20260311113522.mp4", "463"), ("Cam2", "", "172"), ("5", "", "")]
    assert parse_track_lines("") == []


# ---- 順位付け ----
def _entry(cam, tid, name, root=Path("g")):
    d = root / cam / "v" / tid
    return GalleryEntry(face_path=d / f"{name}_face.jpg", body_path=None, track_dir=d, camera_id=cam, video_name="v",
                        track_id=tid, ts=name, video_file="v.mp4")


def test_compute_search_ranks():
    # 比較対象: Cam1 追跡ID1 (g0, g1) / Cam1 追跡ID2 (g2) / Cam2 追跡ID3 (g3) / Cam2 追跡ID4 (g4)
    gallery = [_entry("Cam1", "1", "a"), _entry("Cam1", "1", "b"), _entry("Cam1", "2", "c"),
               _entry("Cam2", "3", "d"), _entry("Cam2", "4", "e")]
    g_emb = np.eye(5, dtype=np.float32)   # 類似度 = q_emb の各成分（検証しやすいように）
    q0 = QueryImage(1, "A", Path("s/A/q0.jpg"), gallery[0], gallery[0].face_path, 0)  # 元画像 g0（追跡ID1）
    q1 = QueryImage(2, "B", Path("s/B/q1.jpg"), gallery[3], gallery[3].face_path, 3)  # 元画像 g3（追跡ID3）
    q2 = QueryImage(3, "", Path("s/q2.jpg"), _entry("-", "-", "q2"), None, None)        # 元画像なし
    q_emb = np.array([[1.0, 0.9, 0.95, 0.8, 0.5],
                      [0.7, 0.6, 0.3, 1.0, 0.2],
                      [0.1, 0.2, 0.3, 0.4, 0.5]], np.float32)
    res = compute_search(q_emb, g_emb, gallery, [q0, q1, q2])
    assert res.round_robin.shape == (3, 3) and np.allclose(res.round_robin, res.round_robin.T)
    assert res.n_tracks == 4 and [(c, t) for _, c, _, t in res.own_tracks] == [("Cam1", "1"), ("Cam2", "3")]
    k1, k3 = res.own_tracks[0][0], res.own_tracks[1][0]
    # q0 → q1（g3 = 0.8）: 全体では g2 0.95, g1 0.9 が上（自分自身 g0 は除く）→ 3 位。Cam2 の中では 1 位
    ir = res.image_ranks[(0, 1)]
    assert (round(ir.sim, 2), ir.rank_camera, ir.rank_all) == (0.8, 1, 3)
    # q0 の追跡ID単位: 追跡ID1 = 0.9（g0 を除いた最大）→ 追跡ID2 0.95 が上で 2 位、他人だけでも 2 位
    t = res.track_ranks[(0, k1)]
    assert (round(t.sim, 2), t.rank_camera, t.rank_all, t.rank_strangers) == (0.9, 2, 2, 2)
    t = res.track_ranks[(0, k3)]   # 追跡ID3 = 0.8: 全体 3 位（0.95, 0.9 が上）、Cam2 では 1 位、他人だけ 2 位（追跡ID2 のみ上）
    assert (t.rank_camera, t.rank_all, t.rank_strangers) == (1, 3, 2)
    # q1 → q0（g0 = 0.7）: 自分自身 g3 を除くと最上位 → 1 位。追跡ID3 は自分の画像しか無いので順位なし
    assert (res.image_ranks[(1, 0)].rank_all, res.image_ranks[(1, 0)].rank_camera) == (1, 1)
    assert (1, k3) not in res.track_ranks and res.track_ranks[(1, k1)].rank_all == 1
    # q2（元画像なし）: 比較対象から何も外さない。q0 の画像 0.1 は 5 位、q1 の画像 0.4 は 2 位
    assert res.image_ranks[(2, 0)].rank_all == 5 and res.image_ranks[(2, 1)].rank_all == 2
    assert (2, 0) in res.image_ranks and (0, 2) not in res.image_ranks   # 元画像の無い q2 を検索対象にした順位は無い
    assert res.track_ranks[(2, k1)].rank_strangers == 3                   # 0.2 より上の他人: 追跡ID2 0.3, 追跡ID4 0.5
    # 本人の追跡IDを追加（Cam2 追跡ID4）→ 他人から外れる。比較対象に無い指定は unmatched に残る
    res2 = compute_search(q_emb, g_emb, gallery, [q0, q1, q2], own_specs=[("Cam2", "v.mp4", "4"), ("Cam9", "", "9")])
    assert [(c, t) for _, c, _, t in res2.own_tracks] == [("Cam1", "1"), ("Cam2", "3"), ("Cam2", "4")]
    assert res2.unmatched_specs == [("Cam9", "", "9")]
    assert res2.track_ranks[(2, k1)].rank_strangers == 2                  # 追跡ID4 を数えなくなる
    # 同じ元画像を 2 か所に置いた（gallery_index が同じ）: 検索元自身を除いた画像なので、互いの画像の順位は出さない
    q0b = QueryImage(4, "C", Path("s/C/q0.jpg"), gallery[0], gallery[0].face_path, 0)
    res3 = compute_search(np.vstack([q_emb, q_emb[:1]]), g_emb, gallery, [q0, q1, q2, q0b])
    assert (0, 3) not in res3.image_ranks and (3, 0) not in res3.image_ranks and (3, 1) in res3.image_ranks
    assert res3.round_robin[0, 3] == pytest.approx(float(q_emb[0] @ q_emb[0]))
    # 比較対象が空でも落ちない
    empty = compute_search(q_emb[:, :0], np.zeros((0, 0), np.float32), [], [q2])
    assert empty.n_tracks == 0 and empty.image_ranks == {} and empty.track_ranks == {}


def test_rank_fill_thresholds():
    assert rank_fill(1).fgColor.rgb == rank_fill(10).fgColor.rgb != rank_fill(11).fgColor.rgb
    assert rank_fill(100).fgColor.rgb != rank_fill(101).fgColor.rgb and rank_fill("-") is None


# ---- Excel・一気通貫 ----
class _MeanExtractor:
    """画素平均を特徴量にする偽のモデル（重み不要）。"""
    embed_dim = 3
    input_size = 16
    needs_landmarks = False
    description = "テスト用"

    def weights_ready(self):
        return True

    def load(self):
        pass

    def embed_batch(self, faces, landmarks=None):
        e = np.stack([f.reshape(-1, 3).mean(axis=0) for f in faces]).astype(np.float32)
        return e / np.maximum(np.linalg.norm(e, axis=1, keepdims=True), 1e-6)


def test_run_query_set_search_end_to_end(tmp_path, monkeypatch):
    import compare_face.core.report.query_set as qs
    monkeypatch.setattr(qs, "get_extractor", lambda name: _MeanExtractor())
    faces = tmp_path / "faces"
    _make_track(faces, "r", "Cam1", "v1", 1, ["20260311_122723_000", "20260311_122724_000"])
    _make_track(faces, "r", "Cam1", "v1", 2, ["20260311_122725_000"])
    d3 = _make_track(faces, "r", "Cam2", "v2", 3, ["20260311_122726_000", "20260311_122727_000"])
    _make_track(faces, "r", "Cam2", "v2", 4, ["20260311_122728_000"], gate_ok=False, gate_reason="切り出しが小さい（短辺 40px < 60px）",
                face_visible=True, crop_min_side=40.0)
    (d3 / "old").mkdir()
    (d3 / "20260311_122727_000_face.jpg").rename(d3 / "old" / "20260311_122727_000_face.jpg")
    # 検索元セット: 追跡ID1 の 2 枚（グループ A）、追跡ID3 の old の 1 枚（グループ B）、比較対象に無い画像 1 枚
    s = create_query_set(default_query_sets_root(faces), "テスト")
    (s / "A").mkdir()
    (s / "B").mkdir()
    shutil.copy2(faces / "r/Cam1/v1/1/20260311_122723_000_face.jpg", s / "A" / "01_正面.jpg")
    shutil.copy2(faces / "r/Cam1/v1/1/20260311_122724_000_face.jpg", s / "A" / "02_横顔.jpg")
    shutil.copy2(d3 / "old" / "20260311_122727_000_face.jpg", s / "B" / "old_の画像.jpg")
    assert imwrite(s / "元画像なし.jpg", _img(4242))
    out, log = run_query_set_search(s, faces, tmp_path / "rep", "AdaFace", "なし", own_specs=[("Cam2", "", "4"), ("Cam9", "", "9")],
                                    echo=lambda m: None)
    assert out.is_file() and out.name.startswith("query_set_")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["条件", "検索元", "総当たり（表）", "総当たり（一覧）", "全体検索（画像）", "全体検索（追跡ID）", "ログ"]
    cond = {r[0]: r[1] for r in wb["条件"].iter_rows(values_only=True)}
    assert cond["検索元 画像数"] == 4 and cond["検索元 元画像を特定した画像数"] == 3 and cond["検索元 比較対象に含まれる画像数"] == 2
    assert cond["比較対象 画像数"] == 4 and cond["比較対象 採用基準外の画像数（除外）"] == 1   # 追跡ID4 は採用基準外、old は含めない
    # 追跡ID3 は検索元の画像が old でも他の画像が比較対象にあるので本人。追跡ID4 は採用基準外で比較対象に無いので追加指定しても入らない
    assert cond["本人の追跡ID（検索元の元画像の追跡ID + 追加指定）"] == "Cam1 / v1.mp4 / 1\nCam2 / v2.mp4 / 3"
    assert "Cam9" in cond["追加指定のうち比較対象に無かったもの"] and "Cam2" in cond["追加指定のうち比較対象に無かったもの"]
    rows = list(wb["検索元"].iter_rows(min_row=4, values_only=True))
    assert [r[0] for r in rows] == [1, 2, 3, 4] and [r[1] for r in rows] == ["-", "A", "A", "B"]
    treat = {r[2]: r[13] for r in rows}
    assert treat["01_正面.jpg"] == "含む" and treat["元画像なし.jpg"].startswith("元画像なし")
    assert treat["old_の画像.jpg"].startswith("old フォルダ") and rows[3][4] == "Cam2" and rows[3][6] == "3"
    pairs = list(wb["総当たり（一覧）"].iter_rows(min_row=4, values_only=True))
    assert len(pairs) == 4 * 3 and all(isinstance(r[9], (int, float)) for r in pairs)  # 1.0 は int で読み戻る
    same = [r for r in pairs if r[10] == "○"]
    assert {(r[1], r[5]) for r in same} == {(2, 3), (3, 2)}   # グループ A の 2 枚どうしだけ
    ws = wb["総当たり（表）"]
    assert ws["A4"].value == 1 and ws["D4"].value == "-" and isinstance(ws["E4"].value, (int, float))
    img = list(wb["全体検索（画像）"].iter_rows(min_row=5, values_only=True))
    assert len(img) == 12
    by = {(r[1], r[4]): r for r in img}
    assert isinstance(by[(2, 3)][9], int) and by[(2, 3)][10] is not None          # 比較対象にある画像どうしは順位あり
    assert by[(2, 1)][8] == "比較対象に無い" and by[(2, 4)][8] == "比較対象に無い"   # 元画像なし / old の画像は順位なし
    trk = wb["全体検索（追跡ID）"]
    assert trk["D4"].value.startswith("Cam1 1") and trk["D5"].value == "類似度" and trk["G5"].value == "他人だけ"
    assert trk["A6"].value == 1 and isinstance(trk["F6"].value, int)   # 検索元 No1 の本人追跡ID（= 自分の追跡ID）の全体順位
    assert any("Excel保存" in (r.name or "") for r in log.records)


def test_set_outside_query_sets_dir_is_not_counted_as_gallery(tmp_path, monkeypatch):
    """CLI で --set を「検索元」以外の場所（比較対象フォルダの中）に置いても、コピーを比較対象に数えず、元画像も正しく見つける。"""
    import compare_face.core.report.query_set as qs
    monkeypatch.setattr(qs, "get_extractor", lambda name: _MeanExtractor())
    faces = tmp_path / "faces"
    _make_track(faces, "r", "Cam1", "v1", 1, ["20260311_122723_000", "20260311_122724_000"])
    s = faces / "a_mysets" / "S"   # 「r」より前に並ぶ名前（コピーの方が先に見つかっても元画像を取る）
    s.mkdir(parents=True)
    shutil.copy2(faces / "r/Cam1/v1/1/20260311_122723_000_face.jpg", s / "20260311_122723_000_face.jpg")
    out, _ = run_query_set_search(s, faces, tmp_path / "rep", "AdaFace", "なし", echo=lambda m: None)
    cond = {r[0]: r[1] for r in openpyxl.load_workbook(out)["条件"].iter_rows(values_only=True)}
    assert cond["比較対象 画像数"] == 2 and cond["検索元 元画像を特定した画像数"] == 1 and cond["検索元 比較対象に含まれる画像数"] == 1
    assert cond["本人の追跡ID（検索元の元画像の追跡ID + 追加指定）"] == "Cam1 / v1.mp4 / 1"


def test_build_report_with_no_own_tracks(tmp_path):
    """元画像の無い検索元だけでも Excel が作れる（本人の追跡IDなし・順位なし）。"""
    s = tmp_path / "s"
    s.mkdir()
    paths = []
    for i in range(2):
        p = s / f"{i}.jpg"
        cv2.imwrite(str(p), _img(i))
        paths.append(p)
    queries = [QueryImage(i + 1, "", p, GalleryEntry(face_path=p, body_path=None, track_dir=s, camera_id="-", video_name="-",
                                                    track_id="-", ts=p.stem), None, None, "元画像なし") for i, p in enumerate(paths)]
    gallery = [_entry("Cam1", "1", "a", tmp_path)]
    cv2.imwrite(str(gallery[0].face_path.parent.mkdir(parents=True, exist_ok=True) or gallery[0].face_path), _img(7))
    res = compute_search(np.array([[1, 0], [0, 1]], np.float32), np.array([[1, 0]], np.float32), gallery, queries)
    out = tmp_path / "o" / "x.xlsx"
    build_query_set_report(out, QuerySetConditions("M", "L", "なし", str(s), str(tmp_path), str(out.parent), datetime.now()),
                           queries, gallery, res, PhaseLogger(echo=lambda m: None))
    wb = openpyxl.load_workbook(out)
    assert wb["全体検索（追跡ID）"]["A6"].value == 1 and wb["全体検索（追跡ID）"].max_column == 3
    assert {r[0]: r[1] for r in wb["条件"].iter_rows(values_only=True)}["本人の追跡ID（検索元の元画像の追跡ID + 追加指定）"] == "なし"
