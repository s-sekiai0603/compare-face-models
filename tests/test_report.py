import json
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
import openpyxl
import pytest

from compare_face.core.features import EXTRACTORS, WeightsUnavailable
from compare_face.core.features.adaface import apply_clahe, preprocess
from compare_face.core.report.gallery import entry_seconds, format_ts, load_entry, scan_gallery, scan_track_folder
from compare_face.core.landmarks import ALIGNERS, get_aligner
from compare_face.core.landmarks.insightface_2d106 import (
    ARCFACE_TEMPLATE_5PTS, align_from_5pts, extract_5pts_from_106, points_within, template_5pts,
)
from compare_face.core.report.merge import assign_groups, merge_embeddings
from compare_face.core.common.phase_log import PhaseLogger
from compare_face.core.report.excel import (
    EXCLUDED_DIR_MESSAGE, EXCLUDED_FILL, GATE_FILL, INVALID_FILL, MODE_IMAGE, QUERY_FILL, SELF_QUERY_MESSAGE, TARGET_FILL,
    ReportConditions, build_report, is_target, search_results,
)
from compare_face.core.common.settings import load_settings
from compare_face.core.report.similarity import aggregate_by_track, competition_ranks, rank_pairs, similarity_matrix
from compare_face.core.extract.validity import face_box_mismatch_reason, foreign_skeleton_reason, revalidate_json
from compare_face.core.report.gate import GateFilter, apply_gate_filter
from compare_face.core.report.quality import QueryFilter, apply_query_filter


def _make_track(root, run, cam, video, tid, stamps, location="本社", body=True, **fields):
    """fields は各画像の JSON 項目の上書き・追加（gate_ok 等）。"""
    d = root / run / cam / video / str(tid)
    d.mkdir(parents=True)
    images = []
    for ts in stamps:
        cv2.imwrite(str(d / f"{ts}_face.jpg"), np.full((90, 80, 3), 100, np.uint8))
        if body:
            cv2.imwrite(str(d / f"{ts}_body.jpg"), np.full((160, 60, 3), 120, np.uint8))
        images.append({"file_face": f"{ts}_face.jpg", "file_body": f"{ts}_body.jpg" if body else None,
                       "person_conf": 0.9, "face_conf": 0.8, "face_width": 42.0, "eye_ratio": 0.31,
                       "frontal": True, "quality": 5.0, "bbox_xyxy": [0, 0, 60, 160],
                       "face_bbox_xyxy": [10, 5, 52, 50], "time_sec": 1.0, **fields})
    (d / f"{tid}.json").write_text(json.dumps(
        {"track_id": tid, "video_file": f"{video}.mp4", "location": location, "images": images},
        ensure_ascii=False), encoding="utf-8")
    return d


def _cond(**kw):
    base = dict(model_name="AdaFace", landmark_name="なし", mode=MODE_IMAGE, query_path="q", gallery_dir="g",
                output_dir="o", show_body=True, generated_at=datetime.now())
    base.update(kw)
    return ReportConditions(**base)


_NOLOG = PhaseLogger(echo=lambda m: None)


# ---- モデルレジストリ ----
def test_registry_has_adaface():
    assert "AdaFace" in EXTRACTORS
    ex = EXTRACTORS["AdaFace"]()
    assert ex.arch == "ir_101" and ex.weights_path().name == "adaface_ir101_webface12m.ckpt"
    assert ex.input_size == 112 and ex.embed_dim == 512
    assert "adaface_ir101_webface12m.ckpt" in ex.manual_instructions()


def test_adaface_preprocess_shape_and_range():
    x = preprocess(np.zeros((30, 10, 3), np.uint8))
    bgr = np.zeros((112, 112, 3), np.uint8)
    bgr[..., 0] = 255  # 青
    assert preprocess(bgr)[0].mean() == 1.0 and preprocess(bgr)[2].mean() == -1.0  # BGR のまま（公式 inference.py）
    assert x.shape == (3, 112, 112) and x.dtype == np.float32 and np.all(x >= -1.0) and np.all(x <= 1.0)
    white = preprocess(np.full((112, 112, 3), 255, np.uint8))
    assert white.max() > 0.99
    img = np.random.randint(0, 255, (50, 40, 3), np.uint8)
    out = apply_clahe(img)
    assert out.shape == img.shape and out.dtype == np.uint8


def test_aligner_registry_and_noalign():
    assert "なし" in ALIGNERS and "InsightFace 2d106det" in ALIGNERS
    al = get_aligner("なし")
    assert al.weights_ready()
    assert al.align(np.zeros((50, 70, 3), np.uint8), 112).shape == (112, 112, 3)
    assert al.align(None, 112) is None
    same = np.zeros((112, 112, 3), np.uint8)
    assert al.align(same, 112) is same
    with pytest.raises(KeyError):
        get_aligner("存在しない")


def test_2d106_helpers():
    lmk = np.zeros((106, 2), np.float32)
    lmk[33:43] = [10, 20]
    lmk[87:97] = [40, 20]
    lmk[86] = [25, 35]
    lmk[52] = [15, 50]
    lmk[61] = [35, 50]
    pts = extract_5pts_from_106(lmk)
    assert pts.tolist() == [[10, 20], [40, 20], [25, 35], [15, 50], [35, 50]]
    assert np.allclose(template_5pts(224), ARCFACE_TEMPLATE_5PTS * 2)
    assert points_within(pts, 50, 60) and not points_within(pts, 20, 20)
    img = np.random.randint(0, 255, (112, 112, 3), np.uint8)
    aligned = align_from_5pts(img, ARCFACE_TEMPLATE_5PTS, 112)  # テンプレートと同じ5点 → ほぼ恒等変換
    assert aligned.shape == (112, 112, 3)
    assert np.mean(np.abs(aligned.astype(int) - img.astype(int))) < 1.0
    assert align_from_5pts(img, np.full((5, 2), np.nan, np.float32), 112) is None


def test_2d106_weights_missing_instructions():
    al = ALIGNERS["InsightFace 2d106det"]()
    if al.weights_ready():
        pytest.skip("2d106det.onnx が配置済み")
    assert "2d106det.onnx" in al.manual_instructions()
    with pytest.raises(WeightsUnavailable):
        al.load()


# ---- ギャラリー ----
def test_format_ts():
    assert format_ts("20260311_122723_500") == "2026/03/11 12:27:23.500"
    assert format_ts("frame00000007_000") == "frame00000007_000"


def test_scan_gallery_reads_json_and_path(tmp_path):
    _make_track(tmp_path, "20261008120000", "Cam1", "20260311122723", 3,
                ["20260311_122723_000", "20260311_122723_500"])
    entries = scan_gallery(tmp_path)
    assert len(entries) == 2
    e = entries[0]
    assert (e.camera_id, e.video_name, e.track_id) == ("Cam1", "20260311122723", "3")
    assert (e.location, e.video_file, e.person_conf) == ("本社", "20260311122723.mp4", 0.9)
    assert (e.face_conf, e.face_width, e.eye_ratio, e.frontal, e.quality) == (0.8, 42.0, 0.31, True, 5.0)
    assert e.embed_path.name.endswith("_face.jpg") and e.body_path.name.endswith("_body.jpg")
    # 採用基準の判定値が無い古い抽出結果: 基準を通った画像として扱う
    assert e.gate_ok and e.gate_reason == "" and e.face_visible is None and e.crop_min_side is None and e.used


def test_scan_ignores_old_folder(tmp_path):
    _make_track(tmp_path, "r", "Cam1", "v", 1, ["20260311_122723_000", "20260311_122724_000"])
    d = tmp_path / "r" / "Cam1" / "v" / "1"
    (d / "old").mkdir()
    for suffix in ("_face.jpg", "_body.jpg"):
        (d / f"20260311_122724_000{suffix}").rename(d / "old" / f"20260311_122724_000{suffix}")
    g = scan_gallery(tmp_path)
    assert [e.ts for e in g] == ["20260311_122723_000"] and g[0].track_id == "1"
    assert [e.ts for e in scan_track_folder(d)] == ["20260311_122723_000"]
    # include_excluded=True: old の画像も元の追跡ID（camera/video/track_id は親フォルダ）として含める
    g2 = scan_gallery(tmp_path, include_excluded=True)
    assert [e.ts for e in g2] == ["20260311_122723_000", "20260311_122724_000"]
    old_entry = g2[1]
    assert (old_entry.camera_id, old_entry.video_name, old_entry.track_id) == ("Cam1", "v", "1")
    assert old_entry.track_dir == d and old_entry.body_path == d / "old" / "20260311_122724_000_body.jpg"
    assert old_entry.face_conf == 0.8  # JSON のメタ情報も元の追跡IDから引ける
    assert [e.ts for e in scan_track_folder(d, include_excluded=True)] == ["20260311_122723_000", "20260311_122724_000"]
    # クエリ追跡IDの除外は old の画像にも効く（クエリ画像が比較対象に漏れない）
    assert scan_gallery(tmp_path, exclude_dirs={d}, include_excluded=True) == []
    (d / "OLD2").mkdir()  # old 以外の名前は対象のまま（ただし追跡IDとして扱われる）
    cv2.imwrite(str(d / "OLD2" / "20260311_122725_000_face.jpg"), np.zeros((20, 20, 3), np.uint8))
    assert len(scan_gallery(tmp_path)) == 2


def test_scan_excludes_query(tmp_path):
    d = _make_track(tmp_path, "r", "C", "v", 1, ["20260311_122723_000"])
    _make_track(tmp_path, "r", "C", "v", 2, ["20260311_122724_000"])
    assert len(scan_gallery(tmp_path, exclude_dirs={d})) == 1
    q = load_entry(next(d.glob("*_face.jpg")))
    assert len(scan_gallery(tmp_path, exclude_files={q.face_path})) == 1
    assert len(scan_track_folder(d)) == 1
    with pytest.raises(ValueError):
        load_entry(d / "x_body.jpg")


def test_missing_json_and_body_fall_back(tmp_path):
    d = _make_track(tmp_path, "r", "C", "v", 1, ["20260311_122723_000"], body=False)
    (d / "1.json").unlink()
    e = scan_gallery(tmp_path)[0]
    assert (e.location, e.person_conf, e.face_conf, e.video_file) == ("-", None, None, "v.mp4")
    assert e.body_path is None


def test_entry_seconds(tmp_path):
    _make_track(tmp_path, "r", "C", "v", 1, ["20260311_122723_000", "20260311_122724_500"])
    e = scan_gallery(tmp_path)
    assert entry_seconds(e[1]) - entry_seconds(e[0]) == 1.5


# ---- 類似度・マージ・ログ ----
def test_similarity_and_aggregation():
    q = np.array([[1, 0]], np.float32)
    g = np.array([[1, 0], [0.6, 0.8], [0, 1], [0.8, 0.6]], np.float32)
    sim = similarity_matrix(q, g)
    pairs = rank_pairs(sim)
    assert [p.g_idx for p in pairs] == [0, 3, 1, 2]
    agg = aggregate_by_track(sim, ["a", "a", "b", "b"])
    assert [(p.g_idx, round(p.sim, 2)) for p in agg] == [(0, 1.0), (3, 0.8)]


def test_competition_ranks():
    assert competition_ranks([0.9, 0.9, 0.8, 0.70001, 0.70004]) == [1, 1, 3, 4, 4]


def test_limit_applies_and_setting_defaults():
    sim = np.arange(10, dtype=np.float32).reshape(1, 10) / 10
    assert len(rank_pairs(sim, 3)) == 3
    assert len(aggregate_by_track(sim, list(range(10)), 4)) == 4
    s = load_settings(path=Path("none.json"))
    assert (s["report_limit"], s["report_model"], s["report_landmark"], s["report_show_body"]) == (500, "AdaFace", "なし", True)
    assert s["report_include_old"] is False
    from compare_face.core.extract.face_crop import MIN_CROP_SIDE, MIN_EYE_RATIO
    assert (s["report_face_gate"], s["report_min_eye_ratio"], s["report_min_crop_side"]) == (True, MIN_EYE_RATIO, MIN_CROP_SIDE)  # 抽出時と同じ基準


def test_assign_groups_fixed_windows_and_skip_empty():
    assert assign_groups([0.0, 0.5, 1.9, 4.0, 5.9, 6.0], 2.0) == [1, 1, 1, 2, 2, 3]
    assert assign_groups([], 1.0) == []


def test_merge_embeddings_mean_and_normalized():
    emb = np.array([[1, 0], [0, 1], [0, 1]], np.float32)
    m = merge_embeddings(emb, [1, 1, 2])
    assert m.shape == (2, 2)
    assert np.allclose(np.linalg.norm(m, axis=1), 1.0)
    assert np.allclose(m[0], [2 ** -0.5, 2 ** -0.5]) and np.allclose(m[1], [0, 1])


def test_phase_logger_records():
    log = PhaseLogger(echo=lambda m: None)
    with log.phase("a") as r:
        r.note = "n"
    try:
        with log.phase("b"):
            raise ValueError("x")
    except ValueError:
        pass
    assert [r.name for r in log.records] == ["a", "b"]
    assert "失敗" in log.records[1].note and log.records[1].end


# ---- Excel ----
def test_build_report_sheets_with_body(tmp_path):
    _make_track(tmp_path, "r", "C", "v", 1, ["20260311_122723_000"])
    _make_track(tmp_path, "r", "C", "v", 2, ["20260311_122724_000", "20260311_122725_000"])
    gallery = scan_gallery(tmp_path)
    queries = gallery[:1]
    sim = np.array([[1.0, 0.5, 0.4]], np.float32)
    out = tmp_path / "rep" / "r.xlsx"
    build_report(out, _cond(gallery_dir=str(tmp_path)), queries, gallery, rank_pairs(sim),
                 aggregate_by_track(sim, [g.track_key for g in gallery]), PhaseLogger(echo=lambda m: None),
                 n_align_failed=2)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["条件", "顔比較", "顔比較(追跡ID集約)", "ログ"]
    ws = wb["顔比較"]
    assert "クエリ" in ws["A1"].value
    text = ws["A5"].value  # 全身表示ON: クエリは3行(1-3)、空行(4)、ギャラリーは5行目から
    assert "1位" in text and "Loc:本社" in text and "顔信頼度:0.800" in text and "顔幅:42px" in text
    assert "正面度:0.31(正面)" in text and "品質:5.0" in text and "映像:v.mp4" in text
    assert len(ws._images) == (1 + 3) * 2  # 顔画像 + 全身画像
    names = [r[1] for r in wb["ログ"].iter_rows(min_row=3, values_only=True)]
    assert any("Excel保存" in n for n in names)
    rows = {r[0]: r[1] for r in wb["条件"].iter_rows(values_only=True)}
    assert rows["顔特徴量抽出モデル"] == "AdaFace" and rows["顔ランドマークモデル"] == "なし"
    assert rows["全身画像の表示"] == "ON" and rows["アライン失敗数（リサイズで代替）"] == 2


def test_build_report_without_body(tmp_path):
    _make_track(tmp_path, "r", "C", "v", 1, ["20260311_122723_000"])
    _make_track(tmp_path, "r", "C", "v", 2, ["20260311_122724_000"])
    g = scan_gallery(tmp_path)
    sim = np.array([[1.0, 0.5]], np.float32)
    out = tmp_path / "nb.xlsx"
    build_report(out, _cond(show_body=False), g[:1], g, rank_pairs(sim),
                 aggregate_by_track(sim, [x.track_key for x in g]), PhaseLogger(echo=lambda m: None))
    ws = openpyxl.load_workbook(out)["顔比較"]
    assert "1位" in ws["A4"].value and len(ws._images) == 3


def test_is_target_and_pink(tmp_path):
    _make_track(tmp_path, "r", "Cam1", "v", 1, ["20260311_122723_000"])
    _make_track(tmp_path, "r", "Cam2", "v", 2, ["20260311_122724_000"])
    g = scan_gallery(tmp_path)
    assert is_target(g[0], [("Cam1", "v.mp4", "1")])
    assert is_target(g[0], [("Cam1", "v", "")])            # 拡張子なし・空欄は不問
    assert not is_target(g[0], [("Cam2", "v.mp4", "1")])
    assert not is_target(g[0], [("", "", "")]) and not is_target(g[0], [])
    sim = np.array([[1.0, 0.5]], np.float32)
    out = tmp_path / "t.xlsx"
    build_report(out, _cond(show_body=False, targets=[("Cam2", "", "2")]), g[:1], g, rank_pairs(sim),
                 aggregate_by_track(sim, [x.track_key for x in g]), PhaseLogger(echo=lambda m: None))
    wb = openpyxl.load_workbook(out)
    ws = wb["顔比較"]
    # 行1-2: クエリ。行4(情報)/5(画像)に順位1(Cam1)・列B=順位2(Cam2)
    assert ws["A4"].fill.fgColor.rgb != TARGET_FILL.fgColor.rgb
    assert ws["B4"].fill.fgColor.rgb == TARGET_FILL.fgColor.rgb
    assert ws["B5"].fill.fgColor.rgb == TARGET_FILL.fgColor.rgb
    cond_rows = {r[0]: r[1] for r in wb["条件"].iter_rows(values_only=True)}
    assert "Cam2" in cond_rows["探したい人物(ピンク表示)"]
    assert "探したい人物検索結果" in wb.sheetnames
    assert cond_rows["old フォルダの画像"].startswith("除外")
    build_report(out, _cond(show_body=False, include_excluded=True), g[:1], g, rank_pairs(sim),
                 aggregate_by_track(sim, [x.track_key for x in g]), PhaseLogger(echo=lambda m: None))
    cond_rows = {r[0]: r[1] for r in openpyxl.load_workbook(out)["条件"].iter_rows(values_only=True)}
    assert cond_rows["old フォルダの画像"].startswith("含める")
    assert cond_rows["比較対象から追加で除外した追跡IDフォルダ"] == "指定なし"


def test_extra_exclude_dirs_in_report(tmp_path):
    """run_report(exclude_dirs=...) の記録: 条件シートに一覧、探したい人物に当たれば検索結果シートに照合対象外と表示。"""
    _make_track(tmp_path, "r", "Cam1", "v", 1, ["20260311_122723_000"])
    _make_track(tmp_path, "r", "Cam2", "v", 2, ["20260311_122724_000"])
    ex_dir = tmp_path / "r" / "Cam1" / "v" / "1"
    g = scan_gallery(tmp_path, exclude_dirs={ex_dir})
    assert [x.track_id for x in g] == ["2"]
    sim = np.array([[0.7]], np.float32)
    out = tmp_path / "x.xlsx"
    build_report(out, _cond(show_body=False, targets=[("Cam1", "v.mp4", "1"), ("Cam2", "", "2")],
                            extra_exclude_dirs=(str(ex_dir),)),
                 g[:1], g, rank_pairs(sim), aggregate_by_track(sim, [x.track_key for x in g]), _NOLOG)
    wb = openpyxl.load_workbook(out)
    cond_rows = {r[0]: r[1] for r in wb["条件"].iter_rows(values_only=True)}
    assert cond_rows["比較対象から追加で除外した追跡IDフォルダ"] == str(ex_dir)
    rows = list(wb["探したい人物検索結果"].iter_rows(values_only=True))
    assert rows[1][2] == EXCLUDED_DIR_MESSAGE and rows[2][2] == 1


def test_search_result_sheet(tmp_path):
    _make_track(tmp_path, "r", "Cam1", "v", 1, ["20260311_122723_000"])
    _make_track(tmp_path, "r", "Cam2", "v", 2, ["20260311_122724_000", "20260311_122725_000"])
    _make_track(tmp_path, "r", "Cam3", "v", 3, ["20260311_122726_000"])
    g = scan_gallery(tmp_path)  # Cam1(1枚), Cam2(2枚), Cam3(1枚)
    sim = np.array([[0.9, 0.5, 0.6, 0.2]], np.float32)
    keys = [x.track_key for x in g]
    targets = [("Cam2", "v.mp4", "2"), ("Cam3", "", "3"), ("Cam9", "", "9"), ("", "", "")]
    res = search_results(targets, g, aggregate_by_track(sim, keys, limit=500), 500, mode=MODE_IMAGE)
    assert [r["no"] for r in res] == [1, 2, 3]  # 空欄行は数えない
    assert (res[0]["rank"], round(res[0]["sim"], 2), res[0]["entry"].track_id, res[0]["hit"]) == (2, 0.6, "2", "-")
    assert res[1]["rank"] == 3 and res[2]["entry"] is None and "該当する人物がありません" in res[2]["message"]
    # 上位 1 件だけなら Cam2 は「存在するが上位に入っていない」
    res1 = search_results(targets, g, aggregate_by_track(sim, keys, limit=1), 1, mode=MODE_IMAGE)
    assert res1[0]["entry"] is None and "上位1件に入っていません" in res1[0]["message"]
    # クエリ自身（比較対象から除外済み）を指定した場合はその旨を表示
    res_self = search_results([("Cam1", "", "1")], g[1:], aggregate_by_track(sim[:, 1:], keys[1:]), 500,
                              mode=MODE_IMAGE, queries=g[:1])
    assert res_self[0]["entry"] is None and res_self[0]["message"] == SELF_QUERY_MESSAGE
    # 追加で除外した追跡IDフォルダ（クエリ画像の元の追跡ID）を指定した場合はその旨を表示（映像名の拡張子・空欄は不問）
    ex_dir = str(tmp_path / "r" / "Cam1" / "v" / "1")
    res_ex = search_results([("Cam1", "v.mp4", "1"), ("", "", "1")], g[1:], aggregate_by_track(sim[:, 1:], keys[1:]), 500,
                            mode=MODE_IMAGE, excluded_dirs=[ex_dir])
    assert [r["message"] for r in res_ex] == [EXCLUDED_DIR_MESSAGE, EXCLUDED_DIR_MESSAGE]
    res_ex2 = search_results([("Cam2", "", "1")], g[1:], aggregate_by_track(sim[:, 1:], keys[1:]), 500,
                             mode=MODE_IMAGE, excluded_dirs=[ex_dir])
    assert "該当する人物がありません" in res_ex2[0]["message"]
    out = tmp_path / "s.xlsx"
    build_report(out, _cond(show_body=True, targets=targets, limit=500), g[:1], g, rank_pairs(sim),
                 aggregate_by_track(sim, keys), PhaseLogger(echo=lambda m: None))
    ws = openpyxl.load_workbook(out)["探したい人物検索結果"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:4] == ("指定No", "指定(カメラID/映像/追跡ID)", "順位", "類似度") and rows[0][-1] == "全身画像"
    assert rows[1][:5] == (1, "Cam2 / v.mp4 / 2", 2, 0.6, "本社") and rows[1][6] == "2"
    assert rows[3][2].startswith("比較対象フォルダに該当する人物がありません")
    assert ws["C2"].fill.fgColor.rgb == TARGET_FILL.fgColor.rgb and len(ws._images) == 4


def test_report_shows_group_ids(tmp_path):
    _make_track(tmp_path, "r", "C", "v", 1, ["20260311_122723_000", "20260311_122730_000"])
    _make_track(tmp_path, "r", "C", "v", 2, ["20260311_122800_000"])
    allx = scan_gallery(tmp_path)
    queries, gallery = allx[:2], allx[2:]
    sim = np.array([[0.9], [0.1]], np.float32)  # 行 = グループ
    out = tmp_path / "g.xlsx"
    build_report(out, _cond(show_body=False, merge_sec=5.0), queries, gallery, rank_pairs(sim),
                 aggregate_by_track(sim, ["k"]), PhaseLogger(echo=lambda m: None), group_ids=[1, 2])
    wb = openpyxl.load_workbook(out)
    ws = wb["顔比較"]
    assert "クエリグループID:G1" in ws["A1"].value and "クエリグループID:G2" in ws["B1"].value
    assert "ヒットしたクエリグループID:G1" in ws["A4"].value  # 最も類似したのは G1
    rows = {r[0]: r[1] for r in wb["条件"].iter_rows(values_only=True)}
    assert "5" in rows["特徴量マージ単位(秒)"] and "2" in rows["特徴量マージ単位(秒)"]


# ---- 無効画像 ----
def test_foreign_skeleton_reason():
    def person(x0):  # 17点を x0..x0+16, y=10 に並べる
        k = np.zeros((17, 3), np.float32)
        k[:, 0] = np.arange(17) + x0; k[:, 1] = 10; k[:, 2] = 0.9
        return k
    own = np.array([0, 0, 20, 50], np.float32)
    boxes = np.array([own, [100, 0, 120, 50]], np.float32)
    kp = np.stack([person(0), person(100)])
    assert foreign_skeleton_reason(own, boxes, kp, (0, 0, 20, 50)) is None        # 自分自身のみ
    kp2 = np.stack([person(0), person(10)])                                         # 他人物の点が範囲内に入る
    r = foreign_skeleton_reason(own, np.array([own, [10, 0, 30, 50]], np.float32), kp2, (0, 0, 20, 50))
    assert r and "別人の骨格" in r
    assert foreign_skeleton_reason(own, None, None, (0, 0, 20, 50)) is None


def _write_invalid(d, ts, reason="別人の骨格を含む（テスト）"):
    jp = next(d.glob("*.json")); data = json.loads(jp.read_text(encoding="utf-8"))
    for im in data["images"]:
        if im["file_face"].startswith(ts):
            im["invalid"] = True; im["invalid_reason"] = reason
    jp.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")


def test_face_box_mismatch_reason():
    box = [100, 100, 140, 150]  # 幅 40 -> 許容 20px
    kp = np.zeros((17, 3), np.float32)
    kp[0] = [120, 125, 0.9]
    assert face_box_mismatch_reason(kp, box) is None                 # 鼻が枠内
    kp[0] = [158, 125, 0.9]
    assert face_box_mismatch_reason(kp, box) is None                 # 枠の外だが顔幅の半分以内（許容内）
    kp[0] = [170, 125, 0.9]
    r = face_box_mismatch_reason(kp, box)
    assert r and "鼻点が無い" in r and "別人の顔" in r and "30px" in r   # 枠から 30px 離れている
    kp[0] = [170, 125, 0.1]
    assert face_box_mismatch_reason(kp, box) is None                 # 鼻が見えていない -> 判定しない
    dicts = [{"x": 170, "y": 125, "conf": 0.5}] + [{"x": 0, "y": 0, "conf": 0}] * 16  # JSON 形式
    assert face_box_mismatch_reason(dicts, box) is not None
    assert face_box_mismatch_reason(None, box) is None and face_box_mismatch_reason(kp, None) is None


def test_revalidate_json(tmp_path):
    def img(nose_x, invalid=False, reason=""):
        return {"file_face": f"{nose_x}_face.jpg", "face_bbox_xyxy": [100, 100, 140, 150], "invalid": invalid,
                "invalid_reason": reason, "keypoints": [{"x": nose_x, "y": 125, "conf": 0.9}] + [{"x": 0, "y": 0, "conf": 0}] * 16}
    p = tmp_path / "7" / "7.json"
    p.parent.mkdir()
    p.write_text(json.dumps({"track_id": 7, "images": [img(120), img(200), img(200, True, "手動除外"),
                                                       img(120, True, "顔枠に本人の鼻点が無い（旧しきい値）")]}), encoding="utf-8")
    assert revalidate_json(p) == (3, 0)  # 1 枚を新たに無効、旧ルールの 1 枚を解除 -> 増減 0
    data = json.loads(p.read_text(encoding="utf-8"))
    assert [im["invalid"] for im in data["images"]] == [False, True, True, False]
    assert "鼻点" in data["images"][1]["invalid_reason"] and data["images"][2]["invalid_reason"] == "手動除外"
    assert data["images"][3]["invalid_reason"] == ""
    assert revalidate_json(p) == (3, 0)  # 2 回目は変化なし


def test_query_filter_and_apply(tmp_path):
    _make_track(tmp_path, "r", "Cam1", "v", 1, ["20260311_122723_000", "20260311_122724_000", "20260311_122725_000"])
    g = scan_gallery(tmp_path)  # face_conf 0.8 / face_width 42
    g[1].face_conf = 0.5
    g[2].invalid = True
    flt = QueryFilter(min_face_conf=0.65, min_norm=17, min_face_width=30)
    assert flt.enabled() and flt.describe() == "顔信頼度 >= 0.65 / 品質ノルム >= 17 / 顔幅 >= 30px"
    assert not QueryFilter().enabled() and QueryFilter().describe() == "指定なし"
    g[1].gate_excluded_reason = "採用基準外"
    assert apply_query_filter(g, QueryFilter(min_face_conf=0.65), None) == 0 and g[1].excluded_reason == ""  # 採用基準外は対象外
    g[1].gate_excluded_reason = ""
    n = apply_query_filter(g, flt, [20.0, 25.0, 10.0])
    assert n == 1 and g[0].excluded_reason == "" and g[0].quality_norm == 20.0
    assert g[1].excluded_reason == "顔信頼度 0.500 < 0.65" and g[2].excluded_reason == ""  # 無効画像は数えない
    assert QueryFilter(min_norm=17).reason(g[2]) == "品質ノルム 10.0 < 17"
    assert QueryFilter(min_face_width=50).reason(g[0]) == "顔幅 42px < 50px"
    g[0].quality_norm = None
    assert QueryFilter(min_norm=17).reason(g[0]) == "品質ノルム未計算"
    # Excel: 除外したクエリは橙、条件シートにしきい値と件数
    g[0].quality_norm = 20.0
    sim = np.array([[1.0, 0.5, 0.2]], np.float32)
    out = tmp_path / "f.xlsx"
    build_report(out, _cond(show_body=False, query_filter_text=flt.describe()), g, g, rank_pairs(sim),
                 aggregate_by_track(sim, [x.track_key for x in g]), PhaseLogger(echo=lambda m: None), used_queries=g[:1])
    wb = openpyxl.load_workbook(out)
    ws = wb["顔比較"]
    assert ws["A1"].fill.fgColor.rgb != EXCLUDED_FILL.fgColor.rgb and "ノルム:20.0" in ws["A1"].value
    assert ws["B1"].fill.fgColor.rgb == EXCLUDED_FILL.fgColor.rgb and "自動選別で除外" in ws["B1"].value
    assert ws["C1"].fill.fgColor.rgb == INVALID_FILL.fgColor.rgb
    cond_rows = {r[0]: r[1] for r in wb["条件"].iter_rows(values_only=True)}
    assert cond_rows["クエリ自動選別（しきい値）"] == flt.describe()
    assert cond_rows["クエリ自動選別で除外した画像数（比較に不使用）"] == 1


# ---- 顔の採用基準（抽出時の顔ゲート・品質ガードの再判定） ----
def test_gate_fields_and_filter(tmp_path):
    _make_track(tmp_path, "r", "C", "v", 1, ["20260311_122723_000"])  # 古い抽出結果（判定値なし）
    _make_track(tmp_path, "r", "C", "v", 2, ["20260311_122724_000"], gate_ok=False, gate_reason="鼻または目が見えない",
                face_visible=False, crop_min_side=70.0)
    _make_track(tmp_path, "r", "C", "v", 3, ["20260311_122725_000"], gate_ok=False, gate_reason="両目の間隔が狭すぎる（横顔）",
                face_visible=True, crop_min_side=70.0, eye_ratio=0.03)
    _make_track(tmp_path, "r", "C", "v", 4, ["20260311_122726_000"], gate_ok=False, gate_reason="切り出しが小さい（短辺 40px < 60px）",
                face_visible=True, crop_min_side=40.0)
    _make_track(tmp_path, "r", "C", "v", 5, ["20260311_122727_000"], gate_ok=True, gate_reason="", face_visible=True, crop_min_side=70.0)
    g = scan_gallery(tmp_path)
    assert [e.gate_ok for e in g] == [True, False, False, False, True]
    assert g[1].face_visible is False and g[0].face_visible is None and g[3].crop_min_side == 40.0
    # 既定（抽出時と同じ）: 抽出時の gate_ok と同じ結果になる。判定値の無い古い抽出結果は通ったとみなす
    flt = GateFilter()
    assert flt.is_default() and flt.enabled() and flt.describe().startswith("抽出時と同じ（既定）")
    assert apply_gate_filter(g, flt) == 3
    assert [e.used for e in g] == [True, False, False, False, True]
    assert g[1].gate_excluded_reason == "鼻または目が見えない"
    assert g[2].gate_excluded_reason == "両目間隔比 0.03 < 0.05（横顔）"
    assert g[3].gate_excluded_reason == "切り出し短辺 40px < 60px"
    # 緩める: 骨格ルールなし・両目間隔比の判定なし・短辺 30px → 全て使う
    loose = GateFilter(face_visible=False, min_eye_ratio=0.0, min_crop_side=30)
    assert apply_gate_filter(g, loose) == 0 and not loose.is_default()
    assert loose.describe() == "骨格ルールなし / 両目間隔比の判定なし / 切り出し短辺 >= 30px"
    assert GateFilter.accept_all().describe() == "適用しない（抽出した顔を全て使う）" and not GateFilter.accept_all().enabled()
    # 厳しくする: 短辺 80px → 抽出時に採用された画像も外れる（判定値の無い古い抽出結果は残る）
    assert apply_gate_filter(g, GateFilter(min_crop_side=80)) == 4
    assert g[4].gate_excluded_reason == "切り出し短辺 70px < 80px" and g[0].gate_excluded_reason == ""
    # 無効画像は数えない
    g[1].invalid = True
    assert apply_gate_filter(g, GateFilter()) == 2 and g[1].gate_excluded_reason == ""


def test_gate_excluded_query_blue_and_conditions(tmp_path):
    _make_track(tmp_path, "r", "C", "v", 1, ["20260311_122723_000", "20260311_122724_000"])
    _make_track(tmp_path, "r", "C", "v", 2, ["20260311_122800_000"], gate_ok=False,
                gate_reason="両目の間隔が狭すぎる（横顔）", face_visible=True, crop_min_side=70.0, eye_ratio=0.03)
    allx = scan_gallery(tmp_path)
    queries, gallery = allx[:2], allx[2:]
    queries[1].gate_excluded_reason = "切り出し短辺 40px < 60px"
    used = [q for q in queries if q.used]
    sim = np.array([[0.9]], np.float32)
    out = tmp_path / "gate.xlsx"
    text = "骨格ルールなし / 両目間隔比の判定なし / 切り出し短辺 >= 60px"
    build_report(out, _cond(show_body=False, gate_filter_text=text), queries, gallery, rank_pairs(sim),
                 aggregate_by_track(sim, ["k"]), _NOLOG, used_queries=used, n_gate_gallery=3)
    wb = openpyxl.load_workbook(out)
    ws = wb["顔比較"]
    assert ws["A1"].fill.fgColor.rgb == QUERY_FILL.fgColor.rgb
    assert ws["B1"].fill.fgColor.rgb == GATE_FILL.fgColor.rgb and ws["B2"].fill.fgColor.rgb == GATE_FILL.fgColor.rgb
    assert "【採用基準外】特徴量比較に未使用" in ws["B1"].value and "40px < 60px" in ws["B1"].value
    # 基準を緩めて使った比較対象画像には、抽出時の判定と切り出し短辺を表示
    assert "抽出時は採用基準外:両目の間隔が狭すぎる（横顔）" in ws["A4"].value and "切出短辺:70px" in ws["A4"].value
    assert "抽出時は採用基準外" not in ws["A1"].value
    rows = {r[0]: r[1] for r in wb["条件"].iter_rows(values_only=True)}
    assert rows["顔の採用基準（抽出時の顔ゲート・品質ガード。クエリ・比較対象に適用）"] == text
    assert rows["クエリ採用基準外の画像数（比較に不使用）"] == 1
    assert rows["比較対象 採用基準外の画像数（除外）"] == 3


def test_gallery_reads_invalid_flag(tmp_path):
    d = _make_track(tmp_path, "r", "C", "v", 1, ["20260311_122723_000", "20260311_122724_000"])
    _write_invalid(d, "20260311_122724_000")
    e = scan_gallery(tmp_path)
    assert [x.invalid for x in e] == [False, True] and e[1].invalid_reason.startswith("別人")


def test_invalid_query_gray_and_excluded(tmp_path):
    d = _make_track(tmp_path, "r", "C", "v", 1, ["20260311_122723_000", "20260311_122724_000"])
    _write_invalid(d, "20260311_122723_000")
    _make_track(tmp_path, "r", "C", "v", 2, ["20260311_122800_000"])
    allx = scan_gallery(tmp_path)
    queries, gallery = allx[:2], allx[2:]
    used = [q for q in queries if not q.invalid]
    sim = np.array([[0.9]], np.float32)
    out = tmp_path / "i.xlsx"
    build_report(out, _cond(show_body=False, merge_sec=5.0), queries, gallery, rank_pairs(sim),
                 aggregate_by_track(sim, ["k"]), PhaseLogger(echo=lambda m: None),
                 group_ids=[None, 1], used_queries=used)
    wb = openpyxl.load_workbook(out); ws = wb["顔比較"]
    assert "【無効画像】" in ws["A1"].value and "別人の骨格を含む（テスト）" in ws["A1"].value
    assert "グループID" not in ws["A1"].value and "クエリグループID:G1" in ws["B1"].value
    assert ws["A1"].fill.fgColor.rgb == INVALID_FILL.fgColor.rgb
    assert ws["A2"].fill.fgColor.rgb == INVALID_FILL.fgColor.rgb       # 画像セルも
    assert ws["B1"].fill.fgColor.rgb != INVALID_FILL.fgColor.rgb
    rows = {r[0]: r[1] for r in wb["条件"].iter_rows(values_only=True)}
    assert rows["クエリ無効画像数（比較に不使用）"] == 1


def test_cache_variant_separates_key_and_file(tmp_path):
    from compare_face.core.report.cache import EmbeddingCache
    a = EmbeddingCache(tmp_path, "M", "L")
    b = EmbeddingCache(tmp_path, "M", "L", variant="clahe")
    assert a.key != b.key and a.path != b.path
    assert "clahe" in b.path.name and "clahe" in b.key


def test_embed_entries_clahe_changes_input(tmp_path):
    """clahe=True ではアライン後の画像に CLAHE が掛かる（単色でない画像は画素が変わる）。"""
    import cv2
    import numpy as np
    from compare_face.core.report.pipeline import embed_entries
    from compare_face.core.report.gallery import load_entry
    from compare_face.core.landmarks import get_aligner

    class _MeanExtractor:
        embed_dim = 3
        input_size = 16
        needs_landmarks = False

        def embed_batch(self, faces, landmarks=None):
            return np.stack([f.reshape(-1, 3).mean(axis=0) for f in faces]).astype(np.float32)

    grad = np.tile(np.linspace(20, 235, 16, dtype=np.uint8), (16, 1))
    img = np.stack([grad, grad // 2, grad // 3], axis=-1)
    p = tmp_path / "Cam1" / "v" / "1"
    p.mkdir(parents=True)
    cv2.imwrite(str(p / "20260101_000000_000_face.jpg"), img)
    entry = load_entry(p / "20260101_000000_000_face.jpg")
    aligner = get_aligner("なし")
    e0, _ = embed_entries(_MeanExtractor(), aligner, [entry])
    e1, _ = embed_entries(_MeanExtractor(), aligner, [entry], clahe=True)
    assert e0.shape == e1.shape == (1, 3)
    assert not np.allclose(e0, e1)
