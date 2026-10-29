"""顔サイズ別精度検証（core/report/verify_face_size.py, verify_face_size_excel.py）。GPU・モデル重みなしで通る。"""
import json
import re
from datetime import datetime

import numpy as np
import openpyxl
import pytest

from compare_face.core.common.imageio import imwrite
from compare_face.core.common.settings import (
    VERIFY_FACE_SIZE_TRACK_COUNT, load_verify_face_size_settings, save_verify_face_size_settings,
)
from compare_face.core.report.gallery import scan_track_folder
from compare_face.core.report.verify_face_size import (
    collect_entries, report_filename, run_verify_face_size, runs_similarity_report, sort_entries,
)
from compare_face.core.report.verify_face_size_excel import (
    FIRST_COL, FIRST_ROW, INFO_SHEET_NAME, SIM_BINS, SHEET_NAME, sim_fill,
)
from test_query_set import _MeanExtractor


def _track(root, cam, video, tid, faces):
    """faces: [(ts, (幅, 高さ) or None)]。顔 bbox は JSON の face_bbox_xyxy だけに入れる（幅・高さは bbox から補完される）。"""
    d = root / cam / video / str(tid)
    d.mkdir(parents=True, exist_ok=True)
    images = []
    for k, (ts, size) in enumerate(faces):
        rng = np.random.default_rng(abs(hash((cam, tid, ts))) % 10000)
        assert imwrite(d / f"{ts}_face.jpg", rng.integers(0, 255, (40, 36, 3), np.uint8))
        im = {"file_face": f"{ts}_face.jpg", "gate_ok": True}
        if size is not None:
            im["face_bbox_xyxy"] = [10, 5, 10 + size[0], 5 + size[1]]
        images.append(im)
    (d / f"{tid}.json").write_text(json.dumps({"video_file": f"{video}.mp4", "images": images}), encoding="utf-8")
    return d


def test_face_height_and_width_from_json_or_bbox(tmp_path):
    d = _track(tmp_path, "Cam1", "v", 1, [("20260311_000001_000", (50, 70))])
    e = scan_track_folder(d)[0]
    assert (e.face_width, e.face_height) == (50.0, 70.0)
    d2 = _track(tmp_path, "Cam1", "v", 2, [("20260311_000002_000", None)])
    e2 = scan_track_folder(d2)[0]
    assert e2.face_width is None and e2.face_height is None


def test_sort_order(tmp_path):
    a = _track(tmp_path, "Cam2", "v", 10, [("20260311_000001_000", (40, 40)), ("20260311_000002_000", (30, 40))])
    b = _track(tmp_path, "Cam1", "v", 9, [("20260311_000003_000", (40, 40)), ("20260311_000004_000", None), ("20260311_000005_000", (99, 20))])
    names = [e.face_path.name for e in sort_entries(collect_entries([a, b, a]))]  # a の重複指定は 1 回だけ
    assert names == ["20260311_000005_000_face.jpg",   # 高さ 20
                     "20260311_000002_000_face.jpg",   # 高さ 40・幅 30
                     "20260311_000003_000_face.jpg",   # 高さ 40・幅 40・Cam1
                     "20260311_000001_000_face.jpg",   # 高さ 40・幅 40・Cam2
                     "20260311_000004_000_face.jpg"]   # サイズなしは末尾


def test_collect_entries_errors_on_missing_dir(tmp_path):
    with pytest.raises(ValueError):
        collect_entries([tmp_path / "ない"])
    assert collect_entries(["", "  "]) == []


def test_sim_fill_bins():
    assert sim_fill(0.05).fgColor.rgb == SIM_BINS[0][1]
    assert sim_fill(0.2).fgColor.rgb == SIM_BINS[1][1]
    assert sim_fill(0.69).fgColor.rgb == SIM_BINS[5][1]
    assert sim_fill(0.95).fgColor.rgb == SIM_BINS[-1][1]


def test_mode_and_filename():
    assert runs_similarity_report("similarity") and runs_similarity_report("both") and not runs_similarity_report("misidentify")
    assert report_filename(datetime(2026, 3, 11, 9, 5, 7)) == "顔サイズごとの精度検証_20260311_090507.xlsx"


def test_run_end_to_end_layout(tmp_path, monkeypatch):
    import compare_face.core.report.verify_face_size as vfs
    monkeypatch.setattr(vfs, "get_extractor", lambda name: _MeanExtractor())
    a = _track(tmp_path / "f", "Cam1", "v1", 1, [("20260311_000001_000", (60, 80)), ("20260311_000002_000", (30, 40))])
    b = _track(tmp_path / "f", "Cam2", "v2", 2, [("20260311_000003_000", (45, 60))])
    out, _ = run_verify_face_size([a, b], tmp_path / "rep", "AdaFace", "なし", echo=lambda m: None)
    assert re.fullmatch(r"顔サイズごとの精度検証_\d{8}_\d{6}\.xlsx", out.name)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames[:2] == [INFO_SHEET_NAME, SHEET_NAME]
    ws = wb[SHEET_NAME]
    # 高さ昇順: 40 → 60 → 80。横は I 列以降、縦は 12 行目以降
    assert [ws.cell(row=11, column=FIRST_COL + j).value for j in range(3)] == [40.0, 60.0, 80.0]
    assert [ws.cell(row=10, column=FIRST_COL + j).value for j in range(3)] == [30.0, 45.0, 60.0]
    assert [ws.cell(row=7, column=FIRST_COL + j).value for j in range(3)] == ["1", "2", "1"]
    assert [ws.cell(row=FIRST_ROW + i, column=3).value for i in range(3)] == ["1", "2", "1"]          # 縦側 C 列 = 追跡ID
    assert [ws.cell(row=FIRST_ROW + i, column=4).value for i in range(3)] == ["v1.mp4", "v2.mp4", "v1.mp4"]  # D 列 = 映像
    assert [ws.cell(row=FIRST_ROW + i, column=7).value for i in range(3)] == [30.0, 45.0, 60.0]
    assert [ws.cell(row=FIRST_ROW + i, column=8).value for i in range(3)] == [40.0, 60.0, 80.0]
    assert ws.cell(row=9, column=FIRST_COL).value is None and ws.cell(row=FIRST_ROW, column=6).value is None  # 顔の向きは空欄
    assert ws["H9"].value == "顔の向き" and ws["F11"].value == "顔の向き"
    # 類似度は左下半分（縦 i > 横 j）のみ
    for i in range(3):
        for j in range(3):
            v = ws.cell(row=FIRST_ROW + i, column=FIRST_COL + j).value
            assert isinstance(v, (int, float)) == (i > j)
            if i > j:
                assert ws.cell(row=FIRST_ROW + i, column=FIRST_COL + j).fill.fgColor.rgb == sim_fill(v).fgColor.rgb


def test_settings_roundtrip(tmp_path):
    p = tmp_path / "s.json"
    assert load_verify_face_size_settings(p)["track_dirs"] == [""] * VERIFY_FACE_SIZE_TRACK_COUNT
    s = load_verify_face_size_settings(p)
    s.update(output_mode="both", track_dirs=["a", "", "c"], output_dir="o", model="X")
    save_verify_face_size_settings(s, p)
    r = load_verify_face_size_settings(p)
    assert r["output_mode"] == "both" and r["model"] == "X" and r["output_dir"] == "o"
    assert r["track_dirs"][:3] == ["a", "", "c"] and len(r["track_dirs"]) == VERIFY_FACE_SIZE_TRACK_COUNT


# ---- 誤識別となる画像の枚数 ----
def test_count_misidentified_row_based():
    from compare_face.core.report.verify_face_size import count_misidentified
    # 2 次元の単位ベクトル。対象者 t0=0°, t1=30°, t2=90°。非対象者 4 枚
    ang = lambda d: [np.cos(np.radians(d)), np.sin(np.radians(d))]
    t = np.array([ang(0), ang(30), ang(90)], np.float32)
    g = np.array([ang(10), ang(20), ang(50), ang(100)], np.float32)
    counts, sim = count_misidentified(t, g)
    # 行 t0（0°）: 非対象者との類似度は 10°>20°>50°>100°。t1(30°) との類似度 cos30=0.866 より高いのは 10°,20° の 2 件
    assert counts[0, 1] == 2
    # 行 t0 × t2（90°）: cos90=0 より高いのは 10°,20°,50° の 3 件（100° は cos<0）
    assert counts[0, 2] == 3
    # 行 t1（30°）: 非対象者 10°(cos20)=0.94, 20°(cos10)=0.985, 50°(cos20)=0.94, 100°(cos70)=0.34。t0(cos30)=0.866 より高いのは 3 件
    assert counts[1, 0] == 3
    assert np.all(np.diag(counts) == 0)
    assert sim[0, 1] == pytest.approx(np.cos(np.radians(30)))
    # 非対象者が空なら全部 0
    assert count_misidentified(t, np.zeros((0, 2), np.float32))[0].sum() == 0


def test_count_misidentified_ties_are_not_counted():
    from compare_face.core.report.verify_face_size import count_misidentified
    t = np.array([[1.0, 0.0], [0.0, 1.0]], np.float32)
    g = np.array([[0.0, 1.0]], np.float32)  # 行 t0 に対し t1 と同じ類似度 0（同値は数えない）
    counts, _ = count_misidentified(t, g)
    assert counts[0, 1] == 0


def test_count_bins_and_colors():
    from compare_face.core.report.verify_face_size_excel import MAX_COUNT_BIN, count_bin, count_fill
    assert [count_bin(n) for n in (0, 1, 10, 11, 20, 21, 50, 51, 100, 101, 150, 151, 200)] == [0, 1, 1, 2, 2, 3, 3, 4, 4, 5, 5, 6, 6]
    assert count_bin(10**6) == MAX_COUNT_BIN
    assert count_fill(0).fgColor.rgb == "FF63BE7B"
    assert len({count_fill(n).fgColor.rgb for n in (0, 1, 11, 21, 51, 101, 10**6)}) == 7


def test_run_misidentify_and_both(tmp_path, monkeypatch):
    import compare_face.core.report.verify_face_size as vfs
    from compare_face.core.report.verify_face_size_excel import MISIDENTIFY_SHEET_NAME
    monkeypatch.setattr(vfs, "get_extractor", lambda name: _MeanExtractor())
    f = tmp_path / "f"
    a = _track(f, "Cam1", "v1", 1, [("20260311_000001_000", (60, 80)), ("20260311_000002_000", (30, 40))])
    _track(f, "Cam2", "v2", 2, [("20260311_000003_000", (45, 60)), ("20260311_000004_000", (46, 61))])
    _track(f, "Cam3", "v3", 3, [("20260311_000005_000", (47, 62))])
    with pytest.raises(ValueError):  # 誤識別レポートには比較対象フォルダが必要
        run_verify_face_size([a], tmp_path / "rep", "AdaFace", "なし", mode="misidentify", echo=lambda m: None)

    out, _ = run_verify_face_size([a], tmp_path / "rep", "AdaFace", "なし", mode="misidentify", gallery_dir=f, echo=lambda m: None)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames[:2] == [INFO_SHEET_NAME, MISIDENTIFY_SHEET_NAME] and SHEET_NAME not in wb.sheetnames
    ws = wb[MISIDENTIFY_SHEET_NAME]
    # 対象者 2 枚（高さ 40, 80）。対角は空欄、それ以外は整数（非対象者は 3 枚）。左下・右上の両方に出る
    assert ws.cell(row=FIRST_ROW, column=FIRST_COL).value is None
    assert ws.cell(row=FIRST_ROW + 1, column=FIRST_COL + 1).value is None
    for i, j in ((0, 1), (1, 0)):
        v = ws.cell(row=FIRST_ROW + i, column=FIRST_COL + j).value
        assert isinstance(v, int) and 0 <= v <= 3
    txt = out.with_suffix(".txt")
    assert txt.is_file()
    cond = {r[0]: r[1] for r in wb["条件"].iter_rows(values_only=True)}
    assert cond["非対象者の画像数"] == 3 and cond["誤認した画像の一覧（テキスト）"] == txt.name

    out2, _ = run_verify_face_size([a], tmp_path / "rep2", "AdaFace", "なし", mode="both", gallery_dir=f, echo=lambda m: None)
    assert openpyxl.load_workbook(out2).sheetnames[:3] == [INFO_SHEET_NAME, SHEET_NAME, MISIDENTIFY_SHEET_NAME]


def test_misidentify_log_lists_images(tmp_path):
    from compare_face.core.report.verify_face_size import count_misidentified, write_misidentify_log
    t_dir = _track(tmp_path, "Cam1", "v1", 1, [("20260311_000001_000", (60, 80)), ("20260311_000002_000", (30, 40))])
    n_dir = _track(tmp_path, "Cam2", "v2", 2, [("20260311_000003_000", (45, 60)), ("20260311_000004_000", (46, 61))])
    ents = sort_entries(scan_track_folder(t_dir))
    nts = scan_track_folder(n_dir)
    ang = lambda d: [np.cos(np.radians(d)), np.sin(np.radians(d))]
    t = np.array([ang(0), ang(60)], np.float32)
    g = np.array([ang(10), ang(80)], np.float32)
    counts, sim = count_misidentified(t, g)
    assert counts[0, 1] == 1 and counts[1, 0] == 2   # 行0(0°): cos60=0.5 より高いのは 10° の 1 件。行1(60°): cos50, cos20 とも cos60 より高い 2 件
    p = tmp_path / "log.txt"
    assert write_misidentify_log(p, ents, nts, t, g, counts, sim) == 2
    text = p.read_text(encoding="utf-8-sig")
    assert text.count("誤認した非対象者の画像 1 件") == 1 and text.count("誤認した非対象者の画像 2 件") == 1
    assert str(nts[0].face_path) in text and str(nts[1].face_path) in text


def test_info_sheet_layout_formula_images_and_order(tmp_path, monkeypatch):
    import compare_face.core.report.verify_face_size as vfs
    monkeypatch.setattr(vfs, "get_extractor", lambda name: _MeanExtractor())
    f = tmp_path / "f"
    a = _track(f, "Cam2", "v1", 10, [("20260311_000002_000", (30, 40)), ("20260311_000001_000", (60, 80))])
    b = _track(f, "Cam1", "v9", 9, [("20260311_000003_000", (45, 60))])
    c = _track(f, "Cam1", "v2", 9, [("20260311_000004_000", (46, 61))])
    body = a / "20260311_000001_000_body.jpg"
    assert imwrite(body, np.random.default_rng(1).integers(0, 255, (120, 50, 3), np.uint8))  # 縦長の全身画像
    out, _ = run_verify_face_size([a, b, c], tmp_path / "rep", "AdaFace", "なし", echo=lambda m: None)
    ws = openpyxl.load_workbook(out)[INFO_SHEET_NAME]
    heads = [ws.cell(row=2, column=k).value for k in range(2, 12)]
    assert heads[0] == "キー情報" and heads[5] == "顔の向き" and heads[-2:] == ["全身画像", "顔画像"]
    # 並び: カメラID → 追跡ID（数値）→ 映像 → 画像ファイル名
    rows = [[ws.cell(row=r, column=k).value for k in range(3, 7)] for r in range(3, 7)]
    assert rows == [["Cam1", "9", "v2.mp4", "20260311_000004_000_face.jpg"], ["Cam1", "9", "v9.mp4", "20260311_000003_000_face.jpg"],
                    ["Cam2", "10", "v1.mp4", "20260311_000001_000_face.jpg"], ["Cam2", "10", "v1.mp4", "20260311_000002_000_face.jpg"]]
    assert ws["B3"].value == '=C3&"|"&D3&"|"&E3&"|"&F3' and ws["B6"].value == '=C6&"|"&D6&"|"&E6&"|"&F6'
    assert ws["G5"].value is None                                   # 顔の向きは空欄
    assert (ws["H5"].value, ws["I5"].value) == (60.0, 80.0)         # 画像サイズ = 顔 bbox の幅・高さ
    assert ws.auto_filter.ref == "B2:K6"
    assert len(ws._images) == 5   # 顔画像 4 枚 + 全身画像 1 枚（_body.jpg があるのは 1 件）
    for im in ws._images:         # 縦横比を保ってセル（行 3〜6 / 列 J・K）に納まる
        a0, a1 = im.anchor._from, im.anchor.to
        assert a0.col == a1.col and a0.row == a1.row and im.anchor.editAs == "twoCell"
        assert a0.col in (9, 10) and 2 <= a0.row <= 5
    body_anchor = next(im.anchor for im in ws._images if im.anchor._from.col == 9)
    w = body_anchor.to.colOff - body_anchor._from.colOff
    h = body_anchor.to.rowOff - body_anchor._from.rowOff
    assert w / h == pytest.approx(50 / 120, rel=0.02)
